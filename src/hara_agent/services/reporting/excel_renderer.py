from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from hara_agent.contracts import ReportContract, ReportFieldMapping
from hara_agent.models import GuidewordDisposition

if TYPE_CHECKING:
    from hara_agent.workflow.state import HARAState


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": MAIN_NS, "r": REL_NS, "pr": PKG_REL_NS}
ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)
ET.register_namespace("x14", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main")
ET.register_namespace("xm", "http://schemas.microsoft.com/office/excel/2006/main")
ET.register_namespace("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
ET.register_namespace("x14ac", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac")
ET.register_namespace("xr", "http://schemas.microsoft.com/office/spreadsheetml/2014/revision")
ET.register_namespace("xr2", "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2")
ET.register_namespace("xr3", "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3")


HARA_REQUIRED_FIELDS = frozenset({"hara_id"})
SG_REQUIRED_FIELDS = frozenset({"sg_id"})
STATIC_HIERARCHY = (
    "function", "output", "guideword", "malfunction", "hazard", "scenario",
)


@dataclass(frozen=True)
class _TableLayout:
    sheet: str
    start_row: int
    columns: dict[str, int]
    header_rows: tuple[int, ...]

    @property
    def max_column(self) -> int:
        return max(self.columns.values())

    @property
    def watermark_cell(self) -> str:
        row = max(1, min(self.header_rows) - 1)
        return f"{HARAExcelRenderer._column_letter(min(self.columns.values()))}{row}"


class HARAExcelRenderer:
    """Render canonical results through a compiled, template-owned ReportContract."""

    def __init__(
        self,
        report_contract: ReportContract | None = None,
        *,
        template_hash: str | None = None,
        report_schema: object | None = None,
        method_contract: object | None = None,
    ):
        self.report_contract = report_contract
        self.template_hash = template_hash
        self.report_schema = report_schema
        self.method_contract = method_contract

    def render(
        self,
        state: HARAState,
        template_path: str | Path,
        output_path: str | Path,
        draft: bool = False,
        smoke: bool = False,
    ) -> Path:
        if not draft and not smoke and not state.can_publish:
            raise ValueError("HARAState still has pending reviews or errors; formal report is blocked")
        if self.report_schema is not None:
            if self.method_contract is None:
                raise ValueError("Canonical report rendering requires the active MethodContract")
            from .canonical_renderer import HARAReportWorkbookRenderer, style_template_hash
            from .projection import HARAReportProjectionService

            risk_trace, risk_trace_reference = self._load_committed_risk_trace(state)

            view_model = HARAReportProjectionService(self.report_schema).project(
                state, self.method_contract,
                risk_trace=risk_trace,
                risk_trace_reference=risk_trace_reference,
                style_template_hash=style_template_hash(template_path),
            )
            return HARAReportWorkbookRenderer().render(
                view_model, template_path, output_path, self.report_schema
            )
        template = Path(template_path).expanduser().resolve()
        if not template.is_file():
            raise FileNotFoundError(f"Template not found: {template}")
        contract = self.report_contract or self._compile_report_contract(template)
        if self.template_hash and self._sha256(template) != self.template_hash:
            raise ValueError("Renderer template hash does not match the compiled MethodContract")
        hara_layout = self._layout(contract.hara_fields, HARA_REQUIRED_FIELDS, "HARA")
        sg_layout = self._layout(contract.safety_goal_fields, SG_REQUIRED_FIELDS, "Safety Goal")
        self._validate_template(template, (hara_layout, sg_layout))

        output = Path(output_path).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        hara_rows = self._hara_rows(state)
        sg_rows = self._safety_goal_rows(state)
        watermark = (
            "SMOKE TEST — NOT FOR RELEASE / 测试子集，禁止正式发布"
            if smoke else "DRAFT — NOT FOR RELEASE / 待工程评审，禁止正式发布"
        )
        replacements = {
            hara_layout.sheet: (
                hara_layout, hara_rows, self._merge_ranges(hara_rows, hara_layout),
                {hara_layout.watermark_cell: watermark} if draft or smoke else {},
            ),
            sg_layout.sheet: (
                sg_layout, sg_rows, [],
                {sg_layout.watermark_cell: watermark} if draft or smoke else {},
            ),
        }
        self._rewrite_package(template, output, replacements)
        self._verify_reopen(output, hara_layout, len(hara_rows), sg_layout, len(sg_rows))
        return output

    def _load_committed_risk_trace(self, state: HARAState) -> tuple[dict | None, str]:
        from hara_agent.workflow.review_artifacts import ReviewArtifactReader

        review_root = os.getenv("HARA_REVIEW_ARTIFACT_DIR", "runtime/review")
        trace_path = (
            ReviewArtifactReader(state.run_id, review_root).directory
            / "risk_execution_trace.json"
        )
        if not trace_path.is_file():
            if state.risk_results:
                raise ValueError(
                    f"Risk execution trace is required for scored state: {trace_path}"
                )
            return None, ""
        try:
            trace = json.loads(trace_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise ValueError(f"Risk execution trace is malformed: {trace_path}") from exc
        if not isinstance(trace, dict) or trace.get("run_id") != state.run_id:
            raise ValueError(f"Risk execution trace run ID does not match state: {trace_path}")
        method_hash = str(self.method_contract.metadata.get(
            "method_source_hash", self.method_contract.metadata.get("template_hash", "")
        ))
        if method_hash and trace.get("method_contract_hash") != method_hash:
            raise ValueError(f"Risk execution trace method hash does not match state: {trace_path}")
        return trace, str(trace_path)

    @staticmethod
    def _compile_report_contract(template: Path) -> ReportContract:
        from hara_agent.template import TemplateRoleCompiler

        method = TemplateRoleCompiler().compile_method(template)
        if method.blocking_diagnostics:
            raise ValueError(
                "Template ReportContract cannot be compiled: "
                + "; ".join(item.message for item in method.blocking_diagnostics)
            )
        return method.report_contract

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _layout(
        mappings: tuple[ReportFieldMapping, ...],
        required: frozenset[str],
        label: str,
    ) -> _TableLayout:
        if not mappings:
            raise ValueError(f"{label} ReportContract has no field mappings")
        sheets = {item.sheet for item in mappings}
        if len(sheets) != 1:
            raise ValueError(f"{label} fields span multiple sheets: {sorted(sheets)}")
        columns: dict[str, int] = {}
        used_columns: dict[int, str] = {}
        for item in mappings:
            if item.canonical_field in columns:
                raise ValueError(f"Duplicate {label} field mapping: {item.canonical_field}")
            if item.column_index < 1:
                raise ValueError(f"Invalid {label} column index: {item.column_index}")
            if item.column_index in used_columns:
                raise ValueError(
                    f"Duplicate {label} output column {item.column_index}: "
                    f"{used_columns[item.column_index]} and {item.canonical_field}"
                )
            columns[item.canonical_field] = item.column_index
            used_columns[item.column_index] = item.canonical_field
        missing = sorted(required - set(columns))
        if missing:
            raise ValueError(f"{label} ReportContract missing required fields: {missing}")
        header_rows = {row for item in mappings for row in item.header_rows}
        if not header_rows:
            raise ValueError(f"{label} ReportContract has no header rows")
        return _TableLayout(
            next(iter(sheets)), max(header_rows) + 1, columns, tuple(sorted(header_rows))
        )

    @staticmethod
    def _validate_template(template: Path, layouts: tuple[_TableLayout, ...]) -> None:
        workbook = load_workbook(template, read_only=False, data_only=False)
        try:
            for layout in layouts:
                if layout.sheet not in workbook.sheetnames:
                    raise ValueError(f"ReportContract sheet is absent: {layout.sheet}")
                if workbook[layout.sheet].max_row < layout.start_row:
                    raise ValueError(
                        f"Template lacks style baseline row {layout.sheet}!{layout.start_row}"
                    )
        finally:
            workbook.close()

    def _hara_rows(self, state: HARAState) -> list[dict[str, object]]:
        scenarios = self._unique_index(state.scenarios, lambda item: item.scenario_id, "Scenario")
        malfunctions = self._unique_index(
            state.malfunctions, lambda item: str(item.get("malfunction_id", "")), "Malfunction",
        )
        functions = self._unique_index(
            state.functions, lambda item: str(item.get("function_id", "")), "Function",
        )
        goals = self._unique_index(state.safety_goals, lambda item: item.sg_id, "Safety Goal")
        function_order = {
            str(item.get("function_id", "")): index
            for index, item in enumerate(state.functions)
        }
        guideword_order = {
            (
                str(item.get("function_id", "")),
                str(item.get("guideword", "")),
            ): index
            for index, item in enumerate(state.guideword_assessments)
        }
        projected: list[tuple[tuple[int, int, int, int], dict[str, object]]] = []

        for assessment_index, assessment in enumerate(state.guideword_assessments):
            disposition = self._guideword_disposition(assessment)
            if disposition is GuidewordDisposition.DOWNSTREAM_CANDIDATE:
                continue
            function_id = str(assessment.get("function_id", ""))
            if function_id not in functions:
                raise ValueError(
                    f"Renderer GuidewordAssessment Function foreign key missing: {function_id!r}"
                )
            function = functions[function_id]
            guideword = str(assessment.get("guideword", ""))
            rationale = self._brief(str(assessment.get("rationale", "")))
            label = (
                "语义不适用"
                if disposition is GuidewordDisposition.NOT_APPLICABLE
                else "未识别到可信车辆级危害"
            )
            projected.append((
                (
                    function_order.get(function_id, len(function_order)),
                    guideword_order.get((function_id, guideword), assessment_index),
                    0,
                    assessment_index,
                ),
                {
                    "function": function.get("name", ""),
                    "output": function.get("output", ""),
                    "guideword": guideword,
                    "malfunction": f"N/A：{rationale}",
                    "remark": f"N/A（{label}）；未进入下游分析",
                },
            ))

        for risk_index, risk in enumerate(state.risk_results):
            if risk.scenario_id not in scenarios or risk.malfunction_id not in malfunctions:
                raise ValueError(
                    f"Renderer foreign key missing: risk={risk.assessment_id!r}, "
                    f"scenario={risk.scenario_id!r}, malfunction={risk.malfunction_id!r}"
                )
            scenario = scenarios[risk.scenario_id]
            malfunction = malfunctions[risk.malfunction_id]
            function_id = str(malfunction.get("function_id", ""))
            if function_id not in functions:
                raise ValueError(f"Renderer Function foreign key missing: {function_id!r}")
            function = functions[function_id]
            goal = goals.get(risk.safety_goal_id)
            if risk.asil.value in {"A", "B", "C", "D"} and (
                not risk.safety_goal_id or goal is None
            ):
                raise ValueError(
                    f"Non-QM risk lacks Safety Goal foreign key: risk={risk.assessment_id!r}, "
                    f"safety_goal_id={risk.safety_goal_id!r}"
                )
            guideword = str(malfunction.get("guideword", ""))
            projected.append((
                (
                    function_order.get(function_id, len(function_order)),
                    guideword_order.get(
                        (function_id, guideword),
                        len(state.guideword_assessments) + risk_index,
                    ),
                    1,
                    risk_index,
                ),
                {
                "function": function.get("name", ""),
                "output": function.get("output", ""),
                "guideword": guideword,
                "malfunction": malfunction.get("description", ""),
                "hazard": malfunction.get("vehicle_level_hazard", ""),
                "scenario": scenario.situational_description,
                "scenario_detail": scenario.situational_detailing,
                "hazardous_event": "；".join(
                    value for value in (risk.hazardous_event, risk.potential_harm) if value
                ),
                "severity": risk.severity.value,
                "severity_rationale": self._basis(risk.severity),
                "exposure": risk.exposure.value,
                "exposure_method": risk.exposure_tf,
                "exposure_rationale": self._basis(risk.exposure),
                "controllability": risk.controllability.value,
                "controllability_rationale": self._basis(risk.controllability),
                "asil": risk.asil.value,
                "sg_id": goal.sg_id if goal else "",
                "safety_goal": goal.text if goal else "",
                "safe_state": goal.safe_state if goal else "",
                "remark": "",
                },
            ))
        projected.sort(key=lambda item: item[0])
        rows = []
        for offset, (_, row) in enumerate(projected, start=1):
            rows.append({"hara_id": f"HARA_{offset:03d}", **row})
        return rows

    @staticmethod
    def _guideword_disposition(assessment: dict) -> GuidewordDisposition:
        raw = assessment.get("disposition")
        raw_value = getattr(raw, "value", raw)
        if raw_value:
            return GuidewordDisposition(str(raw_value))
        return (
            GuidewordDisposition.DOWNSTREAM_CANDIDATE
            if assessment.get("applicable") is True
            else GuidewordDisposition.NOT_APPLICABLE
        )

    @staticmethod
    def _brief(value: str, limit: int = 120) -> str:
        compact = " ".join(value.split()) or "未提供进一步说明"
        if len(compact) <= limit:
            return compact
        return compact[: limit - 1].rstrip() + "…"

    @staticmethod
    def _unique_index(values, key, label: str):
        result = {}
        for value in values:
            identity = key(value)
            if not identity:
                raise ValueError(f"Renderer {label} ID cannot be empty")
            if identity in result:
                raise ValueError(f"Renderer {label} ID is not unique: {identity!r}")
            result[identity] = value
        return result

    @staticmethod
    def _safety_goal_rows(state: HARAState) -> list[dict[str, object]]:
        malfunctions = {str(item.get("malfunction_id", "")): item for item in state.malfunctions}
        risks_by_goal: dict[str, list] = {}
        for risk in state.risk_results:
            if risk.safety_goal_id:
                risks_by_goal.setdefault(risk.safety_goal_id, []).append(risk)
        rows: list[dict[str, object]] = []
        for offset, goal in enumerate(state.safety_goals):
            hazards = list(dict.fromkeys(
                str(malfunctions.get(risk.malfunction_id, {}).get("vehicle_level_hazard", ""))
                for risk in risks_by_goal.get(goal.sg_id, [])
                if malfunctions.get(risk.malfunction_id, {}).get("vehicle_level_hazard")
            ))
            rows.append({
                "hazard_id": f"HZ_{offset + 1:02d}",
                "hazard": "；".join(hazards),
                "sg_id": goal.sg_id,
                "safety_goal": goal.text,
                "safe_state": goal.safe_state,
                "max_asil": goal.max_asil,
            })
        return rows

    @staticmethod
    def _basis(evidence) -> str:
        if evidence.sources and evidence.sources[0].excerpt:
            return evidence.sources[0].excerpt
        return evidence.review_reason

    def _rewrite_package(self, template: Path, output: Path, replacements: dict) -> None:
        with ZipFile(template, "r") as source:
            sheet_paths = self._sheet_paths(source)
            unknown = sorted(set(replacements) - set(sheet_paths))
            if unknown:
                raise ValueError(f"Template sheet relationships missing: {unknown}")
            changed = {
                sheet_paths[name]: self._replace_sheet_data(source.read(sheet_paths[name]), *config)
                for name, config in replacements.items()
            }
            removed_parts: set[str] = set()
            if "xl/calcChain.xml" in source.namelist():
                removed_parts.add("xl/calcChain.xml")
                changed["xl/_rels/workbook.xml.rels"] = self._remove_calc_chain_relationship(
                    source.read("xl/_rels/workbook.xml.rels")
                )
                changed["[Content_Types].xml"] = self._remove_calc_chain_content_type(
                    source.read("[Content_Types].xml")
                )
            handle, temp_name = tempfile.mkstemp(
                prefix=f".{output.stem}.", suffix=".xlsx", dir=str(output.parent)
            )
            os.close(handle)
            try:
                with ZipFile(temp_name, "w", ZIP_DEFLATED) as target:
                    for info in source.infolist():
                        if info.filename in removed_parts:
                            continue
                        target.writestr(info, changed.get(info.filename, source.read(info.filename)))
                os.replace(temp_name, output)
            finally:
                if os.path.exists(temp_name):
                    os.unlink(temp_name)

    @staticmethod
    def _sheet_paths(package: ZipFile) -> dict[str, str]:
        workbook = ET.fromstring(package.read("xl/workbook.xml"))
        relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        targets = {
            item.attrib["Id"]: item.attrib["Target"].lstrip("/")
            for item in relationships.findall(f"{{{PKG_REL_NS}}}Relationship")
        }
        result = {}
        for sheet in workbook.find(f"{{{MAIN_NS}}}sheets"):
            relationship_id = sheet.attrib[f"{{{REL_NS}}}id"]
            target = targets[relationship_id]
            result[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
        return result

    def _replace_sheet_data(
        self,
        xml_bytes: bytes,
        layout: _TableLayout,
        rows: list[dict[str, object]],
        merge_ranges: list[str],
        header_updates: dict[str, str],
    ) -> bytes:
        root = ET.fromstring(xml_bytes)
        sheet_data = root.find("m:sheetData", NS)
        existing_rows = list(sheet_data.findall("m:row", NS))
        template_row = next(
            (row for row in existing_rows if int(row.attrib["r"]) == layout.start_row), None
        )
        if template_row is None:
            raise ValueError(f"Template lacks style baseline row: {layout.start_row}")
        style_by_column = {
            self._column_number(cell.attrib["r"]): cell.attrib.get("s")
            for cell in template_row.findall("m:c", NS)
        }
        row_attributes = {key: value for key, value in template_row.attrib.items() if key != "r"}
        for row in existing_rows:
            if int(row.attrib["r"]) >= layout.start_row:
                sheet_data.remove(row)
        for offset, values in enumerate(rows):
            row_number = layout.start_row + offset
            row = ET.Element(f"{{{MAIN_NS}}}row", {"r": str(row_number), **row_attributes})
            for field, column in sorted(layout.columns.items(), key=lambda item: item[1]):
                value = values.get(field, "")
                if value is None or value == "":
                    continue
                attrs = {"r": f"{self._column_letter(column)}{row_number}", "t": "inlineStr"}
                if style_by_column.get(column) is not None:
                    attrs["s"] = style_by_column[column]
                cell = ET.SubElement(row, f"{{{MAIN_NS}}}c", attrs)
                inline = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
                text = ET.SubElement(inline, f"{{{MAIN_NS}}}t")
                text.text = str(value)
            sheet_data.append(row)
        for reference, value in header_updates.items():
            self._set_inline_string(sheet_data, reference, value)
        self._replace_data_merges(root, layout.start_row, merge_ranges)
        dimension = root.find("m:dimension", NS)
        if dimension is not None:
            last_row = max(layout.start_row - 1, layout.start_row + len(rows) - 1)
            dimension.attrib["ref"] = f"A1:{self._column_letter(layout.max_column)}{last_row}"
        serialized = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return self._preserve_root_namespaces(xml_bytes, serialized)

    @staticmethod
    def _preserve_root_namespaces(original: bytes, serialized: bytes) -> bytes:
        """Retain template namespace declarations referenced lexically by mc:Ignorable."""
        original_text = original.decode("utf-8")
        serialized_text = serialized.decode("utf-8")
        original_root = re.search(r"<worksheet\b[^>]*>", original_text)
        serialized_root = re.search(r"<worksheet\b[^>]*>", serialized_text)
        if original_root is None or serialized_root is None:
            raise ValueError("Worksheet XML root element is missing")
        declarations = {
            prefix: uri
            for prefix, _quote, uri in re.findall(
                r"\sxmlns:([A-Za-z_][\w.-]*)=([\"'])(.*?)\2",
                original_root.group(0),
            )
        }
        current = set(re.findall(
            r"\sxmlns:([A-Za-z_][\w.-]*)=", serialized_root.group(0)
        ))
        additions = "".join(
            f' xmlns:{prefix}="{uri}"'
            for prefix, uri in declarations.items()
            if prefix not in current
        )
        if not additions:
            return serialized
        insertion = serialized_root.end() - 1
        return (serialized_text[:insertion] + additions + serialized_text[insertion:]).encode("utf-8")

    @staticmethod
    def _remove_calc_chain_relationship(xml_bytes: bytes) -> bytes:
        root = ET.fromstring(xml_bytes)
        for relationship in list(root):
            if relationship.attrib.get("Type", "").endswith("/calcChain"):
                root.remove(relationship)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _remove_calc_chain_content_type(xml_bytes: bytes) -> bytes:
        root = ET.fromstring(xml_bytes)
        for override in list(root):
            if override.attrib.get("PartName") == "/xl/calcChain.xml":
                root.remove(override)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @classmethod
    def _set_inline_string(cls, sheet_data, reference: str, value: str) -> None:
        row_number = cls._range_max_row(reference)
        row = next(
            (item for item in sheet_data.findall("m:row", NS) if int(item.attrib["r"]) == row_number),
            None,
        )
        if row is None:
            row = ET.Element(f"{{{MAIN_NS}}}row", {"r": str(row_number)})
            for index, existing in enumerate(list(sheet_data)):
                if int(existing.attrib["r"]) > row_number:
                    sheet_data.insert(index, row)
                    break
            else:
                sheet_data.append(row)
        cell = next(
            (item for item in row.findall("m:c", NS) if item.attrib.get("r") == reference), None
        )
        if cell is None:
            cell = ET.Element(f"{{{MAIN_NS}}}c", {"r": reference})
            row.append(cell)
        for child in list(cell):
            cell.remove(child)
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
        text = ET.SubElement(inline, f"{{{MAIN_NS}}}t")
        text.text = value

    @staticmethod
    def _replace_data_merges(root, start_row: int, new_ranges: list[str]) -> None:
        merge_cells = root.find("m:mergeCells", NS)
        if merge_cells is None:
            merge_cells = ET.Element(f"{{{MAIN_NS}}}mergeCells")
            sheet_data = root.find("m:sheetData", NS)
            root.insert(list(root).index(sheet_data) + 1, merge_cells)
        for merge in list(merge_cells):
            if HARAExcelRenderer._range_max_row(merge.attrib["ref"]) >= start_row:
                merge_cells.remove(merge)
        for reference in new_ranges:
            ET.SubElement(merge_cells, f"{{{MAIN_NS}}}mergeCell", {"ref": reference})
        merge_cells.attrib["count"] = str(len(list(merge_cells)))

    @staticmethod
    def _merge_ranges(rows: list[dict[str, object]], layout: _TableLayout) -> list[str]:
        ranges = []
        available_hierarchy = tuple(
            field for field in STATIC_HIERARCHY if field in layout.columns
        )
        for position, field in enumerate(available_hierarchy):
            column = layout.columns[field]
            parents = available_hierarchy[:position]
            group_start = 0
            for index in range(1, len(rows) + 1):
                boundary = index == len(rows)
                if not boundary:
                    boundary = not (
                        rows[index].get(field) == rows[index - 1].get(field)
                        and all(
                            rows[index].get(parent) == rows[index - 1].get(parent)
                            for parent in parents
                        )
                    )
                if boundary:
                    if index - group_start > 1 and rows[group_start].get(field) not in (None, ""):
                        letter = HARAExcelRenderer._column_letter(column)
                        ranges.append(
                            f"{letter}{layout.start_row + group_start}:"
                            f"{letter}{layout.start_row + index - 1}"
                        )
                    group_start = index
        return ranges

    @staticmethod
    def _range_max_row(reference: str) -> int:
        end = reference.split(":")[-1]
        return int("".join(character for character in end if character.isdigit()))

    @staticmethod
    def _column_number(reference: str) -> int:
        letters = "".join(character for character in reference if character.isalpha())
        result = 0
        for character in letters:
            result = result * 26 + ord(character.upper()) - 64
        return result

    @staticmethod
    def _column_letter(number: int) -> str:
        result = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            result = chr(65 + remainder) + result
        return result

    def _verify_reopen(
        self,
        output: Path,
        hara: _TableLayout,
        hara_count: int,
        sg: _TableLayout,
        sg_count: int,
    ) -> None:
        with ZipFile(output, "r") as package:
            if package.testzip() is not None:
                raise ValueError("Report ZIP package integrity check failed")
            self._verify_package_xml(package)
        workbook = load_workbook(output, read_only=False, data_only=False)
        try:
            hara_sheet = workbook[hara.sheet]
            for offset in range(hara_count):
                row = hara.start_row + offset
                if not hara_sheet.cell(row, hara.columns["hara_id"]).value:
                    raise ValueError(f"Report verification failed: {hara.sheet} row {row} lacks HARA ID")
                if "asil" in hara.columns and str(
                    hara_sheet.cell(row, hara.columns["asil"]).value or ""
                ).startswith("="):
                    raise ValueError(f"Report verification failed: {hara.sheet} row {row} retains ASIL formula")
            sg_sheet = workbook[sg.sheet]
            for offset in range(sg_count):
                row = sg.start_row + offset
                if not sg_sheet.cell(row, sg.columns["sg_id"]).value:
                    raise ValueError(f"Report verification failed: {sg.sheet} row {row} lacks SG ID")
        finally:
            workbook.close()

    @staticmethod
    def _verify_package_xml(package: ZipFile) -> None:
        names = set(package.namelist())
        if "xl/calcChain.xml" in names:
            raise ValueError("Report retains a stale calculation chain")
        for name in sorted(names):
            if not (name.endswith(".xml") or name.endswith(".rels")):
                continue
            content = package.read(name)
            try:
                ET.fromstring(content)
            except ET.ParseError as exc:
                raise ValueError(f"Report XML is malformed: {name}: {exc}") from exc
            if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                continue
            text = content.decode("utf-8")
            root = re.search(r"<worksheet\b[^>]*>", text)
            if root is None:
                raise ValueError(f"Report worksheet root is missing: {name}")
            ignorable = re.search(r"\bmc:Ignorable=([\"'])(.*?)\1", root.group(0))
            if ignorable is None:
                continue
            declared = set(re.findall(
                r"\sxmlns:([A-Za-z_][\w.-]*)=", root.group(0)
            ))
            missing = sorted(set(ignorable.group(2).split()) - declared)
            if missing:
                raise ValueError(
                    f"Report worksheet has undeclared mc:Ignorable prefixes: "
                    f"{name}: {missing}"
                )
