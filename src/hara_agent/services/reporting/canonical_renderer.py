from __future__ import annotations

import hashlib
import os
import tempfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .report_schema import ReportField, ReportSchema
from .view_model import HARAReportViewModel


MAIN_SHEET = "04_HARA"
SAFETY_GOAL_SHEET = "06_Safety Goal"
HEADER_TOP_ROW = 4
HEADER_BOTTOM_ROW = 5
DATA_START_ROW = 6
LEGACY_REFERENCE_SHEETS = frozenset({
    "AI-process", "04_HAZOP", "Severity", "Exposure", "Controllability",
    "VDA702 Summary", "10_SitKat_alt", "VDA702 Full", "ASIL_Table", "Werte",
    "Scenarios_Library",
})
NEW_REPORT_SHEETS = (
    "00_Summary", "01_Item Definition", "02_Functions", "03_Malfunctions",
    "04A_Scenario Detail", "05_Method Basis", "99_Audit",
)
TECHNICAL_MAIN_FIELDS = frozenset({
    "malfunction_id", "scenario_id", "hazardous_event_id", "function_id",
    "assessment_status", "clarification_ids",
})


@dataclass(frozen=True)
class _ColumnPresentation:
    field: str
    column: int
    body_token: str
    source_column: int


@dataclass(frozen=True)
class RenderedReportRegion:
    """The smallest content-and-merge boundary of a clean report sheet."""

    min_row: int
    max_row: int
    min_column: int
    max_column: int

    @classmethod
    def from_sheet(cls, sheet: Any) -> "RenderedReportRegion":
        cells = [
            (cell.row, cell.column)
            for cell in sheet._cells.values()
            if cell.value not in (None, "")
        ]
        if not cells:
            raise ValueError(f"Rendered report sheet has no content: {sheet.title}")
        for merged_range in sheet.merged_cells.ranges:
            anchor = sheet.cell(merged_range.min_row, merged_range.min_col)
            if anchor.value not in (None, ""):
                cells.extend(((merged_range.min_row, merged_range.min_col), (merged_range.max_row, merged_range.max_col)))
        rows, columns = zip(*cells, strict=True)
        return cls(min(rows), max(rows), min(columns), max(columns))

    @property
    def reference(self) -> str:
        return (
            f"{get_column_letter(self.min_column)}{self.min_row}:"
            f"{get_column_letter(self.max_column)}{self.max_row}"
        )

    def contains_cell(self, row: int, column: int) -> bool:
        return self.min_row <= row <= self.max_row and self.min_column <= column <= self.max_column

    def contains_range(self, min_row: int, max_row: int, min_column: int, max_column: int) -> bool:
        return (
            self.min_row <= min_row <= max_row <= self.max_row
            and self.min_column <= min_column <= max_column <= self.max_column
        )


class TemplateStyleRegistry:
    """Template-owned visual tokens used by the report renderer."""

    def __init__(self, workbook: Any):
        self.workbook = workbook
        self.hara = self._required_sheet("05_HARA")
        self.safety_goals = self._required_sheet("06_Safety Goal")
        self.summary = self._required_sheet("0_Document Information")
        self._styles = {
            "hara_title": self._style(self.hara, "A2"),
            "static_header": self._style(self.hara, "A4"),
            "group_header": self._style(self.hara, "I4"),
            "subheader": self._style(self.hara, "J5"),
            "data_static": self._style(self.hara, "A6"),
            "data_harm": self._style(self.hara, "I6"),
            "data_pending_warning": self._style(self.hara, "J6"),
            "data_rationale": self._style(self.hara, "K6"),
            "data_asil": self._style(self.hara, "Q6"),
            "data_safety_goal": self._style(self.hara, "S6"),
            "data_remark": self._style(self.hara, "U6"),
            "safety_goal_title": self._style(self.safety_goals, "A1"),
            "safety_goal_header": self._style(self.safety_goals, "A3"),
            "safety_goal_data": self._style(self.safety_goals, "A4"),
            "summary_title": self._style(self.summary, "B2"),
            "summary_label": self._style(self.summary, "B4"),
        }
        self._source_widths = {
            column: self.hara.column_dimensions[get_column_letter(column)].width
            for column in range(1, self.hara.max_column + 1)
        }
        self._source_heights = {
            row: self.hara.row_dimensions[row].height
            for row in range(1, max(self.hara.max_row, DATA_START_ROW) + 1)
        }

    def _required_sheet(self, name: str) -> Any:
        if name not in self.workbook.sheetnames:
            raise ValueError(f"Template shell lacks required sheet: {name}")
        return self.workbook[name]

    @staticmethod
    def _style(sheet: Any, coordinate: str) -> Any:
        return copy(sheet[coordinate]._style)

    def apply(self, cell: Any, token: str, *, wrap: bool = False) -> None:
        cell._style = copy(self._styles[token])
        if wrap:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment

    def apply_width(self, sheet: Any, target_column: int, source_column: int) -> None:
        width = self._source_widths.get(source_column)
        if width is not None:
            sheet.column_dimensions[get_column_letter(target_column)].width = width

    def apply_page_setup(self, target: Any, source: Any) -> None:
        target.page_setup = copy(source.page_setup)
        target.page_margins = copy(source.page_margins)
        target.print_options = copy(source.print_options)
        target.sheet_properties = copy(source.sheet_properties)
        target.sheet_view.showGridLines = source.sheet_view.showGridLines

    def template_audit(self) -> dict[str, Any]:
        return {
            "style_source": "HARA_Template_AI_20260327.xlsx",
            "style_mode": "TEMPLATE_SHELL",
            "registered_style_tokens": sorted(self._styles),
            "template_sheet_count": len(self.workbook.sheetnames),
            "template_sheets": list(self.workbook.sheetnames),
            "hara": self._sheet_audit(
                self.hara,
                ("A2", "A4", "I4", "J5", "A6", "I6", "J6", "K6", "Q6", "S6", "U6"),
            ),
            "safety_goal": self._sheet_audit(self.safety_goals, ("A1", "A3", "A4")),
            "summary": self._sheet_audit(self.summary, ("B2", "B4")),
        }

    @staticmethod
    def _sheet_audit(sheet: Any, coordinates: Iterable[str]) -> dict[str, Any]:
        return {
            "name": sheet.title,
            "dimensions": sheet.calculate_dimension(),
            "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
            "freeze_panes": str(sheet.freeze_panes or ""),
            "print_area": str(sheet.print_area or ""),
            "print_title_rows": str(sheet.print_title_rows or ""),
            "page_setup": {
                "orientation": sheet.page_setup.orientation,
                "paper_size": sheet.page_setup.paperSize,
                "fit_to_width": sheet.page_setup.fitToWidth,
                "fit_to_height": sheet.page_setup.fitToHeight,
            },
            "styles": {
                coordinate: {
                    "style_id": sheet[coordinate].style_id,
                    "fill": sheet[coordinate].fill.fill_type,
                    "font_name": sheet[coordinate].font.name,
                    "font_bold": bool(sheet[coordinate].font.bold),
                    "alignment": str(sheet[coordinate].alignment.horizontal or ""),
                }
                for coordinate in coordinates
            },
            "column_widths": {
                get_column_letter(column): sheet.column_dimensions[get_column_letter(column)].width
                for column in range(1, sheet.max_column + 1)
            },
            "conditional_formatting_rule_count": len(sheet.conditional_formatting),
            "data_validation_count": len(sheet.data_validations.dataValidation),
            "data_validation_ranges": [str(item.sqref) for item in sheet.data_validations.dataValidation],
        }


class HARAReportWorkbookRenderer:
    """Render the canonical P2-A projection into the original XLSX shell."""

    def __init__(self) -> None:
        self.last_style_audit: dict[str, Any] = {}
        self.last_new_sheet_style_audit: dict[str, Any] = {}

    def render(
        self,
        view_model: HARAReportViewModel,
        template_path: str | Path,
        output_path: str | Path,
        schema: ReportSchema,
    ) -> Path:
        template = Path(template_path).expanduser().resolve()
        if not template.is_file():
            raise FileNotFoundError(f"Report style template not found: {template}")
        output = Path(output_path).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = load_workbook(template)
        registry = TemplateStyleRegistry(workbook)
        try:
            sheets = self._prepare_workbook_shell(workbook)
            self._render_summary(sheets["00_Summary"], view_model, registry)
            self._render_support_table(
                sheets["01_Item Definition"], "Item Definition", registry,
                (("Source ID", view_model.summary.run_id),
                 ("Status", "SOURCE-LINKED"),
                 ("Note", "Facts are projected from the committed runtime state.")),
            )
            self._render_records(
                sheets["02_Functions"], "Functions", registry,
                ("Function ID", "Function", "Output", "Status"), self._functions(view_model),
            )
            self._render_records(
                sheets["03_Malfunctions"], "Malfunctions", registry,
                ("Malfunction ID", "Function ID", "Guideword", "Malfunction", "Hazard"),
                self._malfunctions(view_model),
            )
            layout = self._render_main_hara(sheets[MAIN_SHEET], view_model, schema, registry)
            self._render_scenario_details(
                sheets["04A_Scenario Detail"], view_model, registry,
            )
            self._render_method_basis(sheets["05_Method Basis"], view_model, registry)
            self._render_safety_goals(sheets[SAFETY_GOAL_SHEET], view_model, registry)
            self._render_audit(sheets["99_Audit"], view_model, registry)
            self._save_atomically(workbook, output)
        finally:
            workbook.close()
        self.verify(
            output, schema, len(view_model.rows), len(view_model.scenario_details),
        )
        self.last_style_audit = self.style_restoration_audit(template, output, layout)
        self.last_new_sheet_style_audit = self.new_sheet_style_isolation_audit(output)
        return output

    @staticmethod
    def audit_template(template_path: str | Path) -> dict[str, Any]:
        workbook = load_workbook(Path(template_path).expanduser().resolve(), read_only=False)
        try:
            return TemplateStyleRegistry(workbook).template_audit()
        finally:
            workbook.close()

    @staticmethod
    def _prepare_workbook_shell(workbook: Any) -> dict[str, Any]:
        hara = workbook["05_HARA"]
        safety_goals = workbook["06_Safety Goal"]
        hara.title = MAIN_SHEET
        safety_goals.title = SAFETY_GOAL_SHEET
        sheets: dict[str, Any] = {
            MAIN_SHEET: hara,
            SAFETY_GOAL_SHEET: safety_goals,
        }
        for name in NEW_REPORT_SHEETS:
            sheets[name] = workbook.create_sheet(title=name)
        visible = [
            sheets["00_Summary"], sheets["01_Item Definition"], sheets["02_Functions"],
            sheets["03_Malfunctions"], sheets[MAIN_SHEET],
            sheets["04A_Scenario Detail"], sheets["05_Method Basis"],
            sheets[SAFETY_GOAL_SHEET], sheets["99_Audit"],
        ]
        for sheet in visible:
            sheet.sheet_state = "visible"
        for sheet in workbook.worksheets:
            if sheet not in visible:
                sheet.sheet_state = "hidden"
        workbook._sheets = visible + [sheet for sheet in workbook.worksheets if sheet not in visible]
        workbook.active = 0
        return sheets

    def _render_summary(self, sheet: Any, view_model: HARAReportViewModel, registry: TemplateStyleRegistry) -> None:
        values = view_model.summary.to_dict()
        scope = values["scope"]

        def shown(value: Any) -> Any:
            if value is None:
                return "NOT RECORDED"
            if isinstance(value, list):
                return ", ".join(str(item) for item in value) or "NONE"
            return value

        rows = [
            ("Run Scope", scope.get("run_scope", "")),
            ("Report Class", scope.get("report_class", "")),
            ("Release Status", values["release_status"]),
            ("Run ID", values["run_id"]),
            ("Scope note", scope.get("scope_note", "")),
            ("Current full-project HARA population", scope.get("full_project_population", "")),
            ("04_HARA row meaning", scope.get("hara_row_meaning", "")),
        ]
        if scope.get("run_scope") == "BOUNDED ENGINEERING SAMPLE":
            rows.extend((label, shown(scope.get(key))) for label, key in (
                ("Provider Attempt Budget", "provider_attempt_budget"),
                ("Sample Function Limit", "sample_function_limit"),
                ("Sample Function IDs", "sample_function_ids"),
                ("Sample Malfunction Limit", "sample_malfunction_limit"),
                ("Sample Malfunction IDs", "sample_malfunction_ids"),
                ("Sample Parent Scenario Limit", "sample_parent_scenario_limit"),
                ("Sample Parent Scenario IDs", "sample_parent_scenario_ids"),
                ("Scenario Pair Limit", "scenario_pair_limit"),
            ))
        rows.extend([
            ("Functions available before Function sampling", shown(scope.get("extracted_function_count"))),
            ("Functions selected / omitted", f"{shown(scope.get('selected_function_count'))} / {shown(scope.get('omitted_function_count'))}"),
            ("Guidewords assessed / applicable / filtered in selected Function scope",
             " / ".join(str(shown(scope.get(key))) for key in (
                 "guideword_assessed_count", "guideword_applicable_count", "guideword_filtered_count"
             ))),
            ("Malfunctions available within selected Function scope", shown(scope.get("generated_malfunction_count_in_selected_function_scope"))),
            ("Malfunctions selected / omitted", f"{shown(scope.get('selected_malfunction_count'))} / {shown(scope.get('omitted_malfunction_count'))}"),
            ("Parent Scenarios available within selected upstream scope", shown(scope.get("available_parent_scenario_count_in_selected_scope"))),
            ("Parent Scenarios selected / omitted", f"{shown(scope.get('selected_parent_scenario_count'))} / {shown(scope.get('omitted_parent_scenario_count'))}"),
            ("Method-instantiated scenario pairs", shown(scope.get("method_instantiated_scenario_pair_count"))),
            ("Analytical children before / after driver branching", f"{shown(scope.get('analytical_child_count_before_driver_branch'))} / {shown(scope.get('analytical_child_count_after_driver_branch'))}"),
            ("Causal retained / excluded / pending", " / ".join(str(shown(scope.get(key))) for key in (
                "causal_retained_count", "causal_excluded_count", "causal_pending_count"
            ))),
            ("Risk scoring invoked / not invoked", f"{shown(scope.get('risk_scored_count'))} / {shown(scope.get('risk_not_invoked_count'))}"),
            ("Complete S/E/C/ASIL chain", shown(scope.get("complete_risk_chain_count"))),
            ("04_HARA review rows", values["eligible_hazardous_event_count"]),
            ("Method source", values["method_source"]),
            ("Method hash", values["method_hash"]), ("Report schema", values["report_schema"]),
            ("Report schema version", values["report_schema_version"]),
            ("Report schema hash", values["report_schema_hash"]),
            ("Style template hash", values["style_template_hash"]),
            ("Report status", values["report_status"]),
            ("Severity finalized / pending", f"{values['severity_finalized']} / {values['severity_pending']}"),
            ("Exposure finalized / pending", f"{values['exposure_finalized']} / {values['exposure_pending']}"),
            ("Controllability finalized / pending", f"{values['controllability_finalized']} / {values['controllability_pending']}"),
            ("ASIL finalized / pending", f"{values['asil_finalized']} / {values['asil_pending']}"),
            ("Engineering clarifications", values["clarification_ids"]),
        ])
        self._render_support_table(sheet, "HARA Engineering Report", registry, rows)
        sheet.column_dimensions["B"].width = 48
        sheet.column_dimensions["C"].width = 85
        for row_number, (label, value) in enumerate(rows, start=5):
            lines = max(
                1,
                (len(str(label)) + 45) // 46,
                (len(str(value or "")) + 81) // 82,
            )
            sheet.row_dimensions[row_number].height = min(120, max(24, 18 * lines))

    def _render_support_table(
        self,
        sheet: Any,
        title: str,
        registry: TemplateStyleRegistry,
        rows: Iterable[tuple[str, Any]],
    ) -> None:
        values = list(rows)
        self._clear_sheet(sheet)
        self._set_title(sheet, "B2:H2", "B2", title, registry, "summary_title")
        self._write_support_headers(sheet, ("Metric", "Value"), registry)
        for row_number, (label, value) in enumerate(values, start=5):
            self._write_cell(sheet, row_number, 2, label, registry, "summary_label", wrap=True)
            self._write_cell(sheet, row_number, 3, value, registry, "data_static", wrap=True)
            sheet.row_dimensions[row_number].height = self._estimated_height((label, value), (28, 70), registry)
        self._finish_support_sheet(sheet, 2, max(5, 4 + len(values)), registry)
        self._clean_new_sheet(sheet)

    def _render_records(
        self,
        sheet: Any,
        title: str,
        registry: TemplateStyleRegistry,
        headers: tuple[str, ...],
        records: list[tuple[Any, ...]],
        width_source_columns: tuple[int, ...] | None = None,
    ) -> None:
        self._clear_sheet(sheet)
        title_last_column = get_column_letter(max(8, 1 + len(headers)))
        self._set_title(
            sheet, f"B2:{title_last_column}2", "B2", title,
            registry, "summary_title",
        )
        self._write_support_headers(sheet, headers, registry)
        for row_number, record in enumerate(records, start=5):
            for offset, value in enumerate(record):
                self._write_cell(sheet, row_number, 2 + offset, value, registry, "data_static", wrap=True)
            sheet.row_dimensions[row_number].height = self._estimated_height(record, (26,) * len(record), registry)
        self._finish_support_sheet(
            sheet, len(headers), max(5, len(records) + 4), registry,
            width_source_columns=width_source_columns,
        )
        self._clean_new_sheet(sheet)

    def _render_scenario_details(
        self, sheet: Any, view_model: HARAReportViewModel,
        registry: TemplateStyleRegistry,
    ) -> None:
        headers = (
            "HARA-ID", "Semantic Group ID", "Variant", "Scenario ID", "Situational Description",
            "Situational Detailing", "Speed Constraint", "Causal Status",
            "Object / Interaction", "Hazardous Event", "Physical Inputs",
            "Driver Branch", "C Decision", "Analysis Basis",
        )
        records = [
            (
                item.hara_id, item.semantic_group_id, item.variant, item.scenario_id,
                item.operational_scenario, item.scenario_detail,
                item.speed_constraint, item.causal_status,
                item.object_interaction_summary, item.hazardous_event,
                item.physical_inputs, item.driver_branch,
                item.controllability_branch, item.analysis_basis,
            )
            for item in view_model.scenario_details
        ]
        self._render_records(
            sheet, "Analytical Scenario Detail", registry, headers, records,
            width_source_columns=(1, 5, 14, 11, 7, 8, 11, 16, 14, 7, 11, 14, 16, 14),
        )

    @staticmethod
    def _clear_sheet(sheet: Any) -> None:
        for merged_range in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged_range))
        for row in sheet.iter_rows():
            for cell in row:
                cell.value = None
        sheet.auto_filter.ref = None
        # openpyxl 3.0 accepts an empty list while its print_area setter
        # rejects None; 3.1 supports both.  The project supports >=3.0.
        sheet.print_area = []

    @staticmethod
    def _set_title(sheet: Any, merged_range: str, coordinate: str, value: str, registry: TemplateStyleRegistry, token: str) -> None:
        sheet.merge_cells(merged_range)
        registry.apply(sheet[coordinate], token, wrap=True)
        sheet[coordinate] = value

    def _write_support_headers(self, sheet: Any, headers: Iterable[str], registry: TemplateStyleRegistry) -> None:
        for offset, header in enumerate(headers):
            self._write_cell(sheet, 4, 2 + offset, header, registry, "subheader", wrap=True)
        sheet.row_dimensions[4].height = registry._source_heights.get(HEADER_BOTTOM_ROW) or 24

    def _finish_support_sheet(
        self, sheet: Any, column_count: int, end_row: int,
        registry: TemplateStyleRegistry,
        *, width_source_columns: tuple[int, ...] | None = None,
    ) -> None:
        registry.apply_page_setup(sheet, registry.summary)
        sources = width_source_columns or tuple(
            min(1 + offset, 21) for offset in range(column_count)
        )
        if len(sources) != column_count:
            raise ValueError("Support-sheet width mapping does not match its columns")
        for offset, source_column in enumerate(sources):
            registry.apply_width(sheet, 2 + offset, source_column)
        last_column = get_column_letter(1 + column_count)
        sheet.freeze_panes = "B5"
        sheet.auto_filter.ref = f"B4:{last_column}{end_row}"
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"B1:{last_column}{end_row}"

    @staticmethod
    def _clean_new_sheet(sheet: Any) -> RenderedReportRegion:
        """Remove any blank formatting outside the derived report boundary."""
        region = RenderedReportRegion.from_sheet(sheet)
        for merged_range in list(sheet.merged_cells.ranges):
            anchor = sheet.cell(merged_range.min_row, merged_range.min_col)
            if (
                anchor.value in (None, "")
                or not region.contains_range(
                    merged_range.min_row, merged_range.max_row,
                    merged_range.min_col, merged_range.max_col,
                )
            ):
                sheet.unmerge_cells(str(merged_range))
        for key, cell in list(sheet._cells.items()):
            if not region.contains_cell(cell.row, cell.column) and cell.value in (None, ""):
                del sheet._cells[key]
        for row, dimension in list(sheet.row_dimensions.items()):
            if not region.min_row <= row <= region.max_row:
                del sheet.row_dimensions[row]
        for column, dimension in list(sheet.column_dimensions.items()):
            column_index = sheet[column + "1"].column
            if not region.min_column <= column_index <= region.max_column:
                del sheet.column_dimensions[column]
        return region

    def _render_main_hara(
        self,
        sheet: Any,
        view_model: HARAReportViewModel,
        schema: ReportSchema,
        registry: TemplateStyleRegistry,
    ) -> tuple[_ColumnPresentation, ...]:
        self._clear_sheet(sheet)
        registry.apply_page_setup(sheet, registry.hara)
        # Legacy template entry controls encode the old method.  The
        # YAML-baseline report is read-only and must not imply those rules are active.
        sheet.data_validations.dataValidation.clear()
        sheet.conditional_formatting._cf_rules.clear()
        presentation = schema.presentation.get("main_hara", {})
        visible_fields = tuple(str(value) for value in presentation.get("visible_fields", ()))
        if not visible_fields:
            raise ValueError("Report schema lacks main_hara.visible_fields")
        if any(field in TECHNICAL_MAIN_FIELDS for field in visible_fields):
            raise ValueError("Technical trace fields cannot occupy the reviewer-facing HARA table")
        fields = {field.canonical_field: field for field in schema.fields_for_sheet(MAIN_SHEET)}
        missing = sorted(set(visible_fields) - set(fields))
        if missing:
            raise ValueError(f"Main HARA presentation references unknown fields: {missing}")
        groups = tuple(presentation.get("groups", ()))
        layout = self._main_layout(visible_fields)
        self._set_title(sheet, f"A2:{get_column_letter(len(layout))}2", "A2", "Hazard Analysis and Risk Assessment (HARA)", registry, "hara_title")
        self._write_main_headers(sheet, fields, layout, groups, registry)
        for presentation_column in layout:
            registry.apply_width(sheet, presentation_column.column, presentation_column.source_column)
        for row_number, row in enumerate(view_model.rows, start=DATA_START_ROW):
            values = row.to_dict()
            for presentation_column in layout:
                field = fields[presentation_column.field]
                self._write_cell(
                    sheet, row_number, presentation_column.column,
                    values.get(field.canonical_field, ""), registry, presentation_column.body_token,
                    wrap=field.wrap_text or field.value_type == "text" or field.display_policy == "human_rationale",
                )
            sheet.row_dimensions[row_number].height = self._estimated_height(
                (values.get(item.field, "") for item in layout),
                tuple(sheet.column_dimensions[get_column_letter(item.column)].width or 15 for item in layout),
                registry,
            )
        last_column = get_column_letter(len(layout))
        last_row = max(HEADER_BOTTOM_ROW, DATA_START_ROW + len(view_model.rows) - 1)
        sheet.freeze_panes = registry.hara.freeze_panes or "A6"
        sheet.auto_filter.ref = f"A5:{last_column}{last_row}"
        sheet.print_title_rows = "1:5"
        sheet.print_area = f"A1:{last_column}{last_row}"
        return layout

    @staticmethod
    def _main_layout(visible_fields: tuple[str, ...]) -> tuple[_ColumnPresentation, ...]:
        style_by_field = {
            "hara_id": ("data_static", 1), "function_name": ("data_static", 2),
            "function_output": ("data_static", 3), "guideword": ("data_static", 4),
            "malfunction": ("data_static", 5), "hazard": ("data_static", 6),
            "operational_scenario": ("data_static", 7), "scenario_detail": ("data_static", 8),
            "hazardous_event": ("data_harm", 9), "potential_harm": ("data_harm", 9),
            "severity": ("data_pending_warning", 10), "severity_rationale": ("data_rationale", 11),
            "exposure": ("data_pending_warning", 12), "exposure_rationale": ("data_rationale", 14),
            "controllability": ("data_pending_warning", 15), "controllability_rationale": ("data_rationale", 16),
            "asil": ("data_asil", 17), "asil_rationale": ("data_rationale", 11),
            "ftti": ("data_asil", 17), "ftti_rationale": ("data_rationale", 11),
            "sg_id": ("data_safety_goal", 18), "safety_goal": ("data_safety_goal", 19),
            "safe_state": ("data_safety_goal", 20), "remark": ("data_remark", 21),
        }
        if set(visible_fields) - set(style_by_field):
            raise ValueError("Main HARA presentation has no template style mapping")
        return tuple(
            _ColumnPresentation(field, index, *style_by_field[field])
            for index, field in enumerate(visible_fields, start=1)
        )

    def _write_main_headers(
        self,
        sheet: Any,
        fields: Mapping[str, ReportField],
        layout: tuple[_ColumnPresentation, ...],
        groups: tuple[Any, ...],
        registry: TemplateStyleRegistry,
    ) -> None:
        layout_by_field = {item.field: item for item in layout}
        group_membership = set()
        for group in groups:
            members = tuple(str(field) for field in group.get("fields", ()))
            if not members:
                continue
            columns = [layout_by_field[field].column for field in members]
            if columns != list(range(min(columns), max(columns) + 1)):
                raise ValueError("Main HARA group fields must be contiguous")
            first, last = min(columns), max(columns)
            anchor = sheet.cell(HEADER_TOP_ROW, first)
            registry.apply(anchor, "group_header", wrap=True)
            anchor.value = str(group["header"])
            if first != last:
                sheet.merge_cells(start_row=HEADER_TOP_ROW, start_column=first, end_row=HEADER_TOP_ROW, end_column=last)
            for field in members:
                self._write_cell(sheet, HEADER_BOTTOM_ROW, layout_by_field[field].column, fields[field].header, registry, "subheader", wrap=True)
            group_membership.update(members)
        for item in layout:
            if item.field in group_membership:
                continue
            anchor = sheet.cell(HEADER_TOP_ROW, item.column)
            registry.apply(anchor, "static_header", wrap=True)
            anchor.value = fields[item.field].header
            sheet.merge_cells(start_row=HEADER_TOP_ROW, start_column=item.column, end_row=HEADER_BOTTOM_ROW, end_column=item.column)
        sheet.row_dimensions[HEADER_TOP_ROW].height = registry._source_heights.get(HEADER_TOP_ROW) or 24
        sheet.row_dimensions[HEADER_BOTTOM_ROW].height = registry._source_heights.get(HEADER_BOTTOM_ROW) or 24

    def _render_method_basis(self, sheet: Any, view_model: HARAReportViewModel, registry: TemplateStyleRegistry) -> None:
        basis = view_model.method_basis.to_dict()
        self._render_records(
            sheet, "Method Basis", registry,
            ("Method source", "Guidewords", "Severity", "Exposure", "Controllability", "ASIL", "FTTI"),
            [tuple(basis[key] for key in ("method_source", "guidewords", "severity", "exposure", "controllability", "asil", "ftti"))],
        )

    def _render_safety_goals(self, sheet: Any, view_model: HARAReportViewModel, registry: TemplateStyleRegistry) -> None:
        self._clear_sheet(sheet)
        registry.apply_page_setup(sheet, registry.safety_goals)
        self._set_title(sheet, "A1:F1", "A1", "Hazards and Safety Goals", registry, "safety_goal_title")
        headers = ("HZ-ID", "Hazard", "SG-ID", "Safety Goal", "Safety State", "Max ASIL")
        for column, header in enumerate(headers, start=1):
            self._write_cell(sheet, 3, column, header, registry, "safety_goal_header", wrap=True)
            registry.apply_width(sheet, column, min(column + 5, 21))
        if view_model.safety_goals:
            for row_number, goal in enumerate(view_model.safety_goals, start=4):
                values = (f"HZ-{row_number - 3:03d}", goal.hazard, goal.sg_id, goal.safety_goal, goal.safe_state, goal.max_asil)
                for column, value in enumerate(values, start=1):
                    self._write_cell(sheet, row_number, column, value, registry, "safety_goal_data", wrap=True)
                sheet.row_dimensions[row_number].height = self._estimated_height(values, (14, 36, 14, 44, 32, 12), registry)
        else:
            self._set_title(sheet, "A4:F4", "A4", "No Safety Goal generated because no ASIL assessment has been finalized.", registry, "safety_goal_data")
            self._set_title(sheet, "A5:F5", "A5", "Pending / Not available — DRAFT — NOT FOR RELEASE", registry, "data_pending_warning")
            sheet.row_dimensions[4].height = self._estimated_height((sheet["A4"].value,), (120,), registry)
            sheet.row_dimensions[5].height = self._estimated_height((sheet["A5"].value,), (120,), registry)
        last_row = max(5, 3 + len(view_model.safety_goals))
        sheet.freeze_panes = registry.safety_goals.freeze_panes or "A4"
        sheet.auto_filter.ref = f"A3:F{last_row}"
        sheet.print_title_rows = "1:3"
        sheet.print_area = f"A1:F{last_row}"

    def _render_audit(self, sheet: Any, view_model: HARAReportViewModel, registry: TemplateStyleRegistry) -> None:
        rows_by_hara_id = {item.hara_id: item for item in view_model.rows}
        references = {item.scenario_id: item for item in view_model.audit_references}
        children = view_model.generated_children
        if children:
            records = []
            for child in children:
                item = references.get(child.scenario_id)
                hara_row = rows_by_hara_id.get(item.hara_id) if item else None
                records.append((
                    item.hara_id if item else "",
                    child.malfunction_id,
                    child.scenario_id,
                    item.hazardous_event_id if item else "",
                    hara_row.function_id if hara_row else "",
                    item.assessment_status if item else child.causal_disposition,
                    item.clarification_ids if item else "",
                    child.semantic_group_id or (item.semantic_group_id if item else ""),
                    child.parent_scenario_id or (item.parent_scenario_id if item else ""),
                    item.variant if item else "",
                    item.selected_atom_ids if item else "",
                    item.source_references if item else "",
                    view_model.method_contract_hash,
                    view_model.schema_hash,
                    view_model.style_template_hash,
                    item.risk_trace_reference if item else (
                        view_model.audit_references[0].risk_trace_reference
                        if view_model.audit_references else ""
                    ),
                    child.driver_branch,
                    child.causal_disposition,
                    child.risk_scoring_invoked,
                ))
        else:
            records = [(
                item.hara_id, rows_by_hara_id[item.hara_id].malfunction_id,
                item.scenario_id, item.hazardous_event_id,
                rows_by_hara_id[item.hara_id].function_id, item.assessment_status,
                item.clarification_ids, item.semantic_group_id,
                item.parent_scenario_id, item.variant, item.selected_atom_ids,
                item.source_references, item.method_hash, item.report_schema_hash,
                item.style_template_hash, item.risk_trace_reference,
                "", "", "",
            ) for item in view_model.audit_references]
        self._render_records(
            sheet, "Audit and Traceability", registry,
            (
                "HARA-ID", "Malfunction ID", "Scenario ID",
                "Hazardous Event ID", "Function ID", "Status",
                "Clarification IDs", "Semantic Group ID", "Parent Scenario ID",
                "Variant", "Selected Atom IDs", "Source References",
                "Method hash", "Report schema hash", "Style template hash",
                "Risk execution trace", "Driver Branch", "Causal Disposition",
                "Risk Scoring Invoked",
            ),
            records,
            width_source_columns=(
                1, 2, 1, 7, 10, 16, 2, 5,
                5, 14, 11, 7, 11, 11, 11, 16, 14, 16, 11,
            ),
        )

    @staticmethod
    def _write_cell(sheet: Any, row: int, column: int, value: Any, registry: TemplateStyleRegistry, token: str, *, wrap: bool) -> None:
        cell = sheet.cell(row, column)
        registry.apply(cell, token, wrap=wrap)
        cell.value = "" if value is None else value

    @staticmethod
    def _estimated_height(values: Iterable[Any], widths: Iterable[float], registry: TemplateStyleRegistry) -> float:
        baseline = registry._source_heights.get(DATA_START_ROW) or 18
        lines = 1
        for value, width in zip(values, widths):
            text = str(value or "")
            capacity = max(10, int(float(width or 14) * 1.15))
            lines = max(lines, min(8, (len(text) + capacity - 1) // capacity))
        return max(baseline, min(180, 15 * lines + 3))

    @staticmethod
    def _functions(view_model: HARAReportViewModel) -> list[tuple[Any, ...]]:
        seen: dict[str, tuple[Any, ...]] = {}
        for row in view_model.rows:
            seen.setdefault(row.function_id, (row.function_id, row.function_name, row.function_output, "SOURCE-LINKED"))
        return list(seen.values())

    @staticmethod
    def _malfunctions(view_model: HARAReportViewModel) -> list[tuple[Any, ...]]:
        seen: dict[str, tuple[Any, ...]] = {}
        for row in view_model.rows:
            seen.setdefault(row.malfunction_id, (row.malfunction_id, row.function_id, row.guideword, row.malfunction, row.hazard))
        return list(seen.values())

    @staticmethod
    def _save_atomically(workbook: Any, output: Path) -> None:
        handle, temporary_name = tempfile.mkstemp(prefix=f".{output.stem}.", suffix=".xlsx", dir=str(output.parent))
        os.close(handle)
        try:
            workbook.save(temporary_name)
            os.replace(temporary_name, output)
        finally:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)

    @staticmethod
    def verify(
        path: Path, schema: ReportSchema, expected_rows: int,
        expected_detail_rows: int | None = None,
    ) -> None:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            required_sheets = {item.name for item in schema.sheets}
            missing = sorted(required_sheets - set(workbook.sheetnames))
            if missing:
                raise ValueError(f"Rendered workbook lacks required report sheets: {missing}")
            hara = workbook[MAIN_SHEET]
            data_rows = [
                row for row in range(DATA_START_ROW, hara.max_row + 1)
                if hara.cell(row, 1).value not in (None, "")
            ]
            if len(data_rows) != expected_rows:
                raise ValueError(f"Rendered HARA row count mismatch: {len(data_rows)} != {expected_rows}")
            if expected_detail_rows is not None:
                detail = workbook["04A_Scenario Detail"]
                detail_rows = [
                    row for row in range(5, detail.max_row + 1)
                    if detail.cell(row, 2).value not in (None, "")
                ]
                if len(detail_rows) != expected_detail_rows:
                    raise ValueError(
                        "Rendered Scenario Detail row count mismatch: "
                        f"{len(detail_rows)} != {expected_detail_rows}"
                    )
            headers = [str(hara.cell(HEADER_BOTTOM_ROW, column).value or hara.cell(HEADER_TOP_ROW, column).value or "") for column in range(1, 25)]
            forbidden = {"Malfunction ID", "Scenario ID", "Hazardous Event ID", "Function ID", "Status", "Clarification ID"}
            if forbidden & set(headers):
                raise ValueError("Technical trace fields leaked into the main HARA table")
            expected_merges = {"A2:X2", "I4:J4", "K4:L4", "M4:N4", "O4:P4", "Q4:R4", "S4:T4", "U4:W4"}
            actual_merges = {str(item) for item in hara.merged_cells.ranges}
            if not expected_merges <= actual_merges:
                raise ValueError("Main HARA grouped headers were not restored")
            values = [
                str(hara.cell(row, column).value or "")
                for row in data_rows for column in range(1, 25)
            ]
            raw_json = sum(value.startswith("{") or value.startswith("[") for value in values)
            leakage = sum(any(token in value for token in ("candidate_atom_ids", "NO_EXPLICIT_METHOD_ALIAS", "FA001")) for value in values)
            if raw_json or leakage:
                raise ValueError(f"Main HARA contains internal projection leakage: raw_json={raw_json}, leakage={leakage}")
            if any(workbook[name].sheet_state != "hidden" for name in LEGACY_REFERENCE_SHEETS if name in workbook.sheetnames):
                raise ValueError("Legacy method/reference sheets must not be active in a YAML-baseline report")
        finally:
            workbook.close()

    @classmethod
    def new_sheet_style_isolation_audit(
        cls,
        output_path: str | Path,
        *,
        before_path: str | Path | None = None,
    ) -> dict[str, Any]:
        """Validate that canonical sheets retain no donor formatting outside content."""
        output = load_workbook(Path(output_path).expanduser().resolve(), read_only=False)
        before = None
        if before_path is not None and Path(before_path).is_file():
            before = load_workbook(Path(before_path).expanduser().resolve(), read_only=False)
        try:
            per_sheet = {name: cls._new_sheet_metrics(output[name]) for name in NEW_REPORT_SHEETS}
            residual_sheets = [
                name for name, metrics in per_sheet.items()
                if metrics["ghost_style_cell_count"]
                or metrics["ghost_border_cell_count"]
                or metrics["ghost_fill_cell_count"]
                or metrics["orphan_merged_range_count"]
                or metrics["orphan_row_style_count"]
                or metrics["orphan_column_style_count"]
            ]
            method_basis = per_sheet["05_Method Basis"]
            if before is not None and "05_Method Basis" in before.sheetnames:
                before_metrics = cls._new_sheet_metrics(before["05_Method Basis"])
                ghost_count_before = before_metrics["ghost_style_cell_count"]
            else:
                ghost_count_before = None
            return {
                "NEW_SHEET_STYLE_ROOT_CAUSE": {
                    "creator": "HARAReportWorkbookRenderer._prepare_workbook_shell",
                    "previous_strategy": "whole-sheet donor clone from the summary shell",
                    "donor_sheet": "0_Document Information",
                    "copied_scope": "entire donor worksheet, including merges, populated-range styles, and dimension metadata",
                    "cause": "Only cell values were cleared before rendering, so historical donor styles extended beyond the declared report region.",
                    "replacement_strategy": "workbook.create_sheet plus TemplateStyleRegistry token application to declared content cells only",
                },
                "sheets_checked": list(NEW_REPORT_SHEETS),
                "per_sheet": per_sheet,
                "05_Method Basis": {
                    "rendered_region": method_basis["rendered_region"],
                    "ghost_count_before": ghost_count_before,
                    "ghost_count_after": method_basis["ghost_style_cell_count"],
                },
                "summary": {
                    "new_sheets_checked": len(NEW_REPORT_SHEETS),
                    "sheets_with_residue": len(residual_sheets),
                    "residual_sheets": residual_sheets,
                    "visual_layout_status": "PASS" if not residual_sheets else "VISUAL_LAYOUT_FAIL",
                },
            }
        finally:
            output.close()
            if before is not None:
                before.close()

    @staticmethod
    def _new_sheet_metrics(sheet: Any) -> dict[str, Any]:
        region = RenderedReportRegion.from_sheet(sheet)
        outside = [
            cell for cell in sheet._cells.values()
            if not region.contains_cell(cell.row, cell.column)
        ]
        styled_outside = [cell for cell in outside if cell.has_style]
        border_outside = [cell for cell in outside if HARAReportWorkbookRenderer._has_border(cell)]
        fill_outside = [cell for cell in outside if cell.fill.fill_type is not None]
        orphan_merges = [
            merged_range for merged_range in sheet.merged_cells.ranges
            if sheet.cell(merged_range.min_row, merged_range.min_col).value in (None, "")
            or not region.contains_range(
                merged_range.min_row, merged_range.max_row,
                merged_range.min_col, merged_range.max_col,
            )
        ]
        orphan_rows = [
            row for row, dimension in sheet.row_dimensions.items()
            if not region.min_row <= row <= region.max_row
            and HARAReportWorkbookRenderer._dimension_has_residue(dimension)
        ]
        orphan_columns = [
            column for column, dimension in sheet.column_dimensions.items()
            if not region.min_column <= sheet[column + "1"].column <= region.max_column
            and HARAReportWorkbookRenderer._dimension_has_residue(dimension)
        ]
        return {
            "rendered_region": region.reference,
            "styled_cells_inside_region": sum(
                cell.has_style for cell in sheet._cells.values()
                if region.contains_cell(cell.row, cell.column)
            ),
            "styled_cells_outside_region": len(styled_outside),
            "ghost_style_cell_count": len(styled_outside),
            "ghost_border_cell_count": len(border_outside),
            "ghost_fill_cell_count": len(fill_outside),
            "orphan_merged_range_count": len(orphan_merges),
            "orphan_merged_ranges": [str(item) for item in orphan_merges],
            "orphan_row_style_count": len(orphan_rows),
            "orphan_column_style_count": len(orphan_columns),
        }

    @staticmethod
    def _has_border(cell: Any) -> bool:
        border = cell.border
        return any(
            getattr(getattr(border, side), "style", None) is not None
            for side in ("left", "right", "top", "bottom", "diagonal")
        )

    @staticmethod
    def _dimension_has_residue(dimension: Any) -> bool:
        return bool(
            dimension.hidden
            or dimension.outlineLevel
            or getattr(dimension, "height", None) is not None
            or getattr(dimension, "width", None) is not None
            or getattr(dimension, "style_id", 0)
        )

    @staticmethod
    def style_restoration_audit(
        template_path: str | Path,
        output_path: str | Path,
        layout: tuple[_ColumnPresentation, ...] | None = None,
    ) -> dict[str, Any]:
        template = load_workbook(Path(template_path).expanduser().resolve(), read_only=False)
        output = load_workbook(Path(output_path).expanduser().resolve(), read_only=False)
        try:
            registry = TemplateStyleRegistry(template)
            hara = output[MAIN_SHEET]
            source = registry.hara
            layout = layout or tuple()
            expected_merges = {"A2:X2", "I4:J4", "K4:L4", "M4:N4", "O4:P4", "Q4:R4", "S4:T4", "U4:W4"}
            title_preserved = hara["A2"].value == source["A2"].value
            header_tokens = {
                "A4": "static_header", "I4": "group_header", "K4": "group_header",
                "I5": "subheader", "K5": "subheader", "Q5": "subheader",
            }
            header_reused = sum(
                HARAReportWorkbookRenderer._visual_style_matches(hara[cell], registry._styles[token])
                for cell, token in header_tokens.items()
            )
            width_reused = sum(
                hara.column_dimensions[get_column_letter(item.column)].width == registry._source_widths.get(item.source_column)
                for item in layout
            )
            body_rows = [
                row for row in range(DATA_START_ROW, hara.max_row + 1)
                if hara.cell(row, 1).value not in (None, "")
            ][:10]
            body_reused = 0
            body_total = 0
            for row in body_rows:
                for item in layout:
                    body_total += 1
                    if HARAReportWorkbookRenderer._visual_style_matches(hara.cell(row, item.column), registry._styles[item.body_token]):
                        body_reused += 1
            headers = [str(hara.cell(HEADER_BOTTOM_ROW, column).value or hara.cell(HEADER_TOP_ROW, column).value or "") for column in range(1, 25)]
            audit_headers = [str(cell.value or "") for cell in output["99_Audit"][4]]
            return {
                "style_source": "HARA_Template_AI_20260327.xlsx",
                "style_mode": "TEMPLATE_SHELL",
                "template_sheet_count": len(template.sheetnames),
                "output_sheet_count": len(output.sheetnames),
                "template_title_preserved": title_preserved,
                "group_structures_preserved": expected_merges <= {str(item) for item in hara.merged_cells.ranges},
                "merged_ranges_preserved_or_remapped": sorted(expected_merges),
                "header_style_reuse": {"matched": header_reused, "total": len(header_tokens), "rate": header_reused / len(header_tokens)},
                "data_style_reuse": {"matched": body_reused, "total": body_total, "rate": body_reused / body_total if body_total else 1.0},
                "column_width_reuse": {"matched": width_reused, "total": len(layout), "rate": width_reused / len(layout) if layout else 1.0},
                "freeze_panes": {
                    "template": str(source.freeze_panes or ""),
                    "output": str(hara.freeze_panes or ""),
                    "status": "PRESERVED" if source.freeze_panes else "ADDED_FOR_REVIEW",
                },
                "print_settings": {
                    "template_orientation": source.page_setup.orientation,
                    "output_orientation": hara.page_setup.orientation,
                    "output_area": str(hara.print_area or ""),
                    "output_titles": str(hara.print_title_rows or ""),
                    "output_filter": str(hara.auto_filter.ref or ""),
                },
                "conditional_formatting": {
                    "template_rule_count": len(source.conditional_formatting),
                    "output_rule_count": len(hara.conditional_formatting),
                    "status": "LEGACY_RULES_EXCLUDED_FROM_YAML_BASELINE_MAIN_HARA",
                },
                "data_validation": {
                    "template_count": len(source.data_validations.dataValidation),
                    "output_count": len(hara.data_validations.dataValidation),
                    "status": "LEGACY_ENTRY_CONTROLS_EXCLUDED_FROM_YAML_BASELINE_MAIN_HARA",
                },
                "new_columns": ["Potential Harm", "FTTI", "FTTI Rationale", "ASIL Rationale"],
                "new_column_style_sources": {
                    "Potential Harm": "05_HARA!I6",
                    "ASIL Rationale": "05_HARA!K6",
                    "FTTI": "05_HARA!Q6",
                    "FTTI Rationale": "05_HARA!K6",
                },
                "technical_fields_moved_from_main": all(value not in headers for value in ("Malfunction ID", "Scenario ID", "Hazardous Event ID", "Function ID", "Status", "Clarification ID")),
                "technical_fields_in_audit": all(value in audit_headers for value in ("Malfunction ID", "Scenario ID", "Hazardous Event ID", "Function ID", "Status", "Clarification IDs")),
                "legacy_reference_sheets_hidden": sorted(name for name in LEGACY_REFERENCE_SHEETS if name in output.sheetnames and output[name].sheet_state == "hidden"),
                "legacy_reference_sheets_excluded_from_visible": all(output[name].sheet_state == "hidden" for name in LEGACY_REFERENCE_SHEETS if name in output.sheetnames),
                "generic_workbook_creation_disabled": True,
                "raw_json_in_main_hara": 0,
            }
        finally:
            template.close()
            output.close()

    @staticmethod
    def _visual_style_matches(cell: Any, expected_style: Any) -> bool:
        return (
            cell._style.fontId == expected_style.fontId
            and cell._style.fillId == expected_style.fillId
            and cell._style.borderId == expected_style.borderId
            and cell._style.numFmtId == expected_style.numFmtId
        )


def style_template_hash(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


__all__ = [
    "HARAReportWorkbookRenderer", "RenderedReportRegion", "TemplateStyleRegistry",
    "style_template_hash",
]
