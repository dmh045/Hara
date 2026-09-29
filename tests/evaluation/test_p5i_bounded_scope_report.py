from __future__ import annotations

from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from hara_agent.services.reporting.offline_rebuild import OfflineReportRebuilder
from hara_agent.services.reporting.projection import HARAReportProjectionService
from hara_agent.workflow.state import HARAState


ROOT = Path(__file__).resolve().parents[2]
CHECKPOINT = ROOT / "runtime/sec_closure/sec_closure_20260929_d.checkpoint.json"
REVIEW = ROOT / "runtime/review/sec_closure_20260929_d/risk_execution_trace.json"
PREVIOUS = ROOT / "output/HARA_Engineering_SEC_Review.xlsx"


def test_unbounded_run_without_risk_execution_does_not_claim_a_measured_population():
    state = HARAState(run_id="full", guideword_assessments=[
        {"function_id": "F1", "status": "FINALIZED", "applicable": True},
        {"function_id": "F2", "status": "FINALIZED", "applicable": True},
        {"function_id": "F2", "status": "FINALIZED", "applicable": False},
    ])
    scope = HARAReportProjectionService._run_scope(state, [], 0)
    assert scope["run_scope"] == "FULL PROJECT HARA"
    assert scope["full_project_population"] == "NOT MEASURED IN THIS RUN"
    assert (scope["guideword_assessed_count"], scope["guideword_applicable_count"],
            scope["guideword_filtered_count"]) == (3, 2, 1)


@pytest.mark.skipif(
    not all(path.is_file() for path in (CHECKPOINT, REVIEW, PREVIOUS)),
    reason="Local SEC closure runtime artifacts are unavailable",
)
def test_current_bounded_run_rebuild_preserves_risk_rows_and_shows_child_lineage(tmp_path):
    output = tmp_path / "scope-fixed.xlsx"
    OfflineReportRebuilder().rebuild(
        checkpoint_path=CHECKPOINT,
        method_baseline_path=ROOT / "method_assets/fusa_baseline_v1/manifest.yaml",
        report_style_template_path=ROOT / "references/HARA_Template_AI_20260327.xlsx",
        output_path=output,
        review_root=ROOT / "runtime/review",
        audit_output_dir=tmp_path / "audits",
    )
    current = load_workbook(output, read_only=True, data_only=True)
    previous = load_workbook(PREVIOUS, read_only=True, data_only=True)
    try:
        summary = {
            row[1]: row[2]
            for row in current["00_Summary"].iter_rows(min_row=5, values_only=True)
            if row[1]
        }
        assert summary["Run Scope"] == "BOUNDED ENGINEERING SAMPLE"
        assert summary["Report Class"] == "ENGINEERING SAMPLE ONLY"
        assert summary["Current full-project HARA population"] == "NOT MEASURED IN THIS RUN"
        assert summary["Functions selected / omitted"] == "1 / 8"
        assert summary["Malfunctions selected / omitted"] == "1 / 6"
        assert summary["Analytical children before / after driver branching"] == "12 / 24"
        assert summary["Causal retained / excluded / pending"] == "13 / 11 / 0"
        assert summary["Risk scoring invoked / not invoked"] == "13 / 11"
        assert summary["04_HARA review rows"] == 13

        current_risks = list(current["04_HARA"].iter_rows(min_row=6, values_only=True))
        previous_risks = list(previous["04_HARA"].iter_rows(min_row=6, values_only=True))
        assert len(current_risks) == len(previous_risks) == 13
        assert [tuple("" if value is None else value for value in row) for row in current_risks] == [
            tuple("" if value is None else value for value in row) for row in previous_risks
        ]
        details = list(current["04A_Scenario Detail"].iter_rows(min_row=5, values_only=True))
        assert len(details) == 13

        audit = current["99_Audit"]
        headers = {cell.value: cell.column - 1 for cell in audit[4] if cell.value}
        inventory = list(audit.iter_rows(min_row=5, values_only=True))
        assert len(inventory) == 24
        assert len({row[headers["Scenario ID"]] for row in inventory}) == 24
        assert Counter(row[headers["Causal Disposition"]] for row in inventory) == {
            "RETAINED_FOR_RISK": 13, "CAUSAL_GAP": 11,
        }
        assert Counter(row[headers["Risk Scoring Invoked"]] for row in inventory) == {
            True: 13, False: 11,
        }
        assert all(row[headers["Parent Scenario ID"]] for row in inventory)
        assert all(row[headers["Driver Branch"]] for row in inventory)
    finally:
        current.close()
        previous.close()
