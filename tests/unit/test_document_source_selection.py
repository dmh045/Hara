from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from docx import Document

from hara_agent.services.extraction import (
    DocumentReader, FactRetrievalSpec, ValidatedArtifactCache,
    select_source_blocks,
)
from hara_agent.models import ItemDefinitionFacts, SourceRef
from hara_agent.services.semantic import ItemEvidenceRouter
from hara_agent.workflow.nodes.extraction import read_item_document
from hara_agent.workflow.nodes.item_artifacts import extract_item_artifacts
from hara_agent.workflow.state import HARAState


ROOT = Path(__file__).resolve().parents[2]
POLICY = {
    "policy_id": "AVP_SEC_2026_09_V1",
    "version": "1",
    "input_selection": {
        "policy_id": "EXCLUDE_NONFUNCTIONAL_PERFORMANCE_SPEED_V1",
        "excluded_sections_for_scoring": ["性能要求"],
        "non_operational_speed_sections": ["执行器能力参数要求"],
    },
}


def _table(document: Document, text: str) -> None:
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "描述"
    table.cell(1, 0).text = "车速"
    table.cell(1, 1).text = text


def test_docx_body_order_and_unstyled_section_provenance(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("功能", style="Heading 2")
    document.add_paragraph("正式功能说明")
    _table(document, "功能输出")
    document.add_paragraph("非功能要求")
    document.add_paragraph("性能要求").runs[0].bold = True
    _table(document, "泊车最高车速≤5km/h")
    document.add_paragraph("执行器能力参数要求").runs[0].bold = True
    _table(document, "制动能力车速范围[-15,15] km/h")
    document.add_paragraph("环境条件").runs[0].bold = True
    _table(document, "控车范围0-7kph")
    path = tmp_path / "ordered.docx"
    document.save(path)

    blocks = DocumentReader().read(path).blocks
    assert [block.block_id for block in blocks] == [
        "P-0001", "P-0002", "T-001-R-0001", "T-001-R-0002",
        "P-0003", "P-0004", "T-002-R-0001", "T-002-R-0002",
        "P-0005", "T-003-R-0001", "T-003-R-0002",
        "P-0006", "T-004-R-0001", "T-004-R-0002",
    ]
    assert blocks[7].section_path[-2:] == ("非功能要求", "性能要求")
    assert blocks[10].section_path[-1] == "执行器能力参数要求"
    assert blocks[13].section_path[-1] == "环境条件"

    selection = select_source_blocks(blocks, POLICY)
    assert {item.block_id: item.reason for item in selection.excluded_blocks} == {
        "T-002-R-0002": "PERFORMANCE_SPEED_EXPECTATION",
        "T-003-R-0002": "ACTUATOR_SPEED_CAPABILITY",
    }
    assert "T-004-R-0002" in {
        block.block_id for block in selection.selected_blocks
    }


def test_real_item_selection_is_shared_by_core_and_targeted_routes() -> None:
    state = read_item_document(
        HARAState("source-selection-test"),
        str(ROOT / "input/ItemDef.docx"),
        project_analysis_policy=POLICY,
    )
    raw = {block["block_id"]: block for block in state.item_definition["raw_blocks"]}
    selected = {block["block_id"]: block for block in state.item_definition["blocks"]}

    assert raw["T-011-R-0002"]["section_path"][-1] == "性能要求"
    assert raw["T-012-R-0003"]["location"] == "table[12].row[3]"
    assert "T-012-R-0003" not in selected
    assert "≤5km/h" not in state.item_definition["text"]
    assert "T-013-R-0011" not in selected
    assert selected["T-014-R-0015"]["section_path"][-1] == "环境条件"
    assert "0-7kph" in selected["T-014-R-0015"]["text"]
    assert "T-014-R-0013" in selected  # Driver allowed set remains available.
    assert "T-006-R-0002" in selected  # Explicit Function membership remains.

    speed_spec = FactRetrievalSpec(
        "speed.operational_context", ("operational speed",),
        ("km/h", "kph"), structural_kind="OPERATIONAL_SPEED_CANDIDATE",
    )
    routed = ItemEvidenceRouter().route(
        state.item_definition["blocks"], "project_evidence",
        required_specs=(speed_spec,),
    )
    assert routed is not None
    assert "T-014-R-0015" in routed.block_ids
    assert "T-012-R-0003" not in routed.block_ids
    assert "T-013-R-0011" not in routed.block_ids
    environmental = next(
        block for block in routed.source_blocks
        if block["block_id"] == "T-014-R-0015"
    )
    assert environmental["section_path"][-1] == "环境条件"


def test_policy_version_changes_extraction_fingerprint_without_docx_change() -> None:
    blocks = DocumentReader().read(ROOT / "input/ItemDef.docx").blocks
    old = select_source_blocks(blocks)
    active = select_source_blocks(blocks, POLICY)
    revised = select_source_blocks(blocks, {**POLICY, "version": "2"})

    assert old.fingerprint != active.fingerprint
    assert active.fingerprint != revised.fingerprint
    assert any(item.block_id == "T-012-R-0003" for item in active.excluded_blocks)
    assert all(item.block_id != "T-012-R-0003" for item in active.selected_blocks)


def test_artifact_cache_misses_when_only_project_policy_version_changes(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("Parking system in a parking area.")
    path = tmp_path / "item.docx"
    document.save(path)

    class CoreAgent:
        PROMPT_VERSION = "core-test-v1"

        def __init__(self) -> None:
            self.client = SimpleNamespace(config=SimpleNamespace(provider="fake", model="fake"))
            self.validator = SimpleNamespace(ensure_valid=lambda functions: None)
            self.calls = 0

        def extract(self, document_text, source_id, blocks, **kwargs):
            self.calls += 1
            return ItemDefinitionFacts(
                "Parking system", "Vehicle control boundary", ["Active"],
                ["parking area"], ["parking road"], ["clear"], ["dry"],
                0, 7, sources=[SourceRef(
                    "item_definition", source_id, "paragraph[1]",
                    "Parking system in a parking area.",
                )],
            ), [], {"llm_call_count": 1, "elapsed_seconds": 0.0}

        @staticmethod
        def _ensure_source_grounded(facts, functions, document_text, blocks):
            return []

    agent = CoreAgent()
    supplement = SimpleNamespace(PROMPT_VERSION="supplement-test-v1")
    cache = ValidatedArtifactCache(tmp_path / "cache", "readwrite")

    def run(policy: dict):
        state = read_item_document(
            HARAState("cache-test"), str(path), project_analysis_policy=policy,
        )
        return extract_item_artifacts(
            state, agent, supplement, ItemEvidenceRouter(), cache=cache,
        )

    first = run(POLICY)
    repeated = run(POLICY)
    revised = run({**POLICY, "version": "2"})

    assert agent.calls == 2
    assert first.item_definition["source_selection"]["fingerprint"] == (
        repeated.item_definition["source_selection"]["fingerprint"]
    )
    assert revised.item_definition["source_selection"]["fingerprint"] != (
        first.item_definition["source_selection"]["fingerprint"]
    )
    assert first.audit_trail[-2]["cache_hit"] is False
    assert repeated.audit_trail[-2]["cache_hit"] is True
    assert revised.audit_trail[-2]["cache_hit"] is False
