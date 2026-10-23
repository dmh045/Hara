"""Independent-review regressions: incomplete group, summary, existing SG intent."""
from collections import Counter
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from hara_agent.models import EvidenceValue, ReviewStatus, SafetyGoal
from hara_agent.services.reporting import HARAReportProjectionService, load_report_schema
from hara_agent.services.reporting.canonical_renderer import HARAReportWorkbookRenderer
from hara_agent.services.reporting.engineering_text_mapper import EngineeringReportTextMapper
from test_reviewer_report_closure import _p6_report_state, _pending
from test_reuse_reporting_acceptance import active_method


def mixed_pending_group():
    state = _p6_report_state()
    source, risk = state.scenarios[0], state.risk_results[0]
    causal = state.item_definition['scenario_assessments'][0]
    state.item_definition['scenario_assessments'] = []
    state.scenarios, state.risk_results = [], []
    for i in range(14):
        child, assessment = deepcopy(source), deepcopy(risk)
        child.scenario_id = assessment.scenario_id = f'SCN-{i}'
        assessment.assessment_id = f'RA-{i}'
        assessment.severity = EvidenceValue('S1', ReviewStatus.FINALIZED, rule_version='fixture')
        assessment.exposure = EvidenceValue('E3' if i < 12 else 'E2', ReviewStatus.FINALIZED, rule_version='fixture')
        assessment.controllability = _pending()
        assessment.asil = _pending()
        state.scenarios.append(child)
        state.risk_results.append(assessment)
        state.item_definition['scenario_assessments'].append({**deepcopy(causal), 'scenario_id': child.scenario_id})
    return state


def test_incomplete_group_does_not_display_first_member_noncommon_e(active_method):
    state = mixed_pending_group()
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert len(view.rows) == 2
    assert {(r.severity, r.exposure, r.controllability, r.asil) for r in view.rows} == {
        ('S1', 'E2', '待评定', '待评定'), ('S1', 'E3', '待评定', '待评定')}
    for row in view.rows:
        assert {d.exposure for d in view.scenario_details if d.hara_id == row.hara_id} == {row.exposure}
        assert '斜线分隔' not in row.remark
    assert len(view.scenario_details) == 14
    assert Counter(d.exposure for d in view.scenario_details) == {'E3': 12, 'E2': 2}
    assert state == before


@pytest.mark.parametrize('field,values', [
    ('driver_in_vehicle', (True, False)),
    ('ego_longitudinal_direction', ('FORWARD', 'REVERSE')),
    ('object_position', ('FRONT', 'REAR')),
    ('ego_heading_relative_to_downhill', ('UPHILL', 'DOWNHILL')),
    ('object_motion_state', ('STATIONARY', 'MOVING')),
])
def test_review_groups_preserve_distinct_conditions_even_when_ratings_are_equal(active_method, field, values):
    state = mixed_pending_group()
    for risk in state.risk_results:
        risk.exposure = EvidenceValue('E3', ReviewStatus.FINALIZED, rule_version='fixture')
    for i, child in enumerate(state.scenarios):
        child.facts[field] = values[i % 2]
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert len(view.rows) == 2
    assert len(view.scenario_details) == 14
    assert {d.scenario_id for d in view.scenario_details} == {s.scenario_id for s in state.scenarios}


def test_same_tuple_atom_and_parameter_variants_share_display_only(active_method):
    state = mixed_pending_group()
    for i, (scenario, risk) in enumerate(zip(state.scenarios, state.risk_results)):
        risk.exposure = EvidenceValue('E3', ReviewStatus.FINALIZED, rule_version='fixture')
        scenario.facts.update(scenario_atom_ids=['SO006' if i % 2 else 'SO010'],
            relative_distance_m=1.001 + i / 10, object_id=f'local-object-{i}')
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert len(view.rows) == 1 and len(view.scenario_details) == 14
    assert view.rows[0].exposure == 'E3' and view.rows[0].controllability == '待评定'
    group = view.projection_metrics['grouping_map'][0]
    assert len(group['scenario_ids']) == len(set(group['scenario_ids'])) == 14
    assert all(key is None for key in group['risk_equivalence_keys'])
    assert {r.selected_atom_ids for r in view.audit_references} == {'SO006', 'SO010'}
    assert len({d.physical_inputs for d in view.scenario_details}) == 14
    assert state == before


def test_same_numeric_tuple_does_not_hide_different_review_states(active_method):
    state = mixed_pending_group()
    state.risk_results[0].severity = EvidenceValue('S1', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert len(view.rows) == 3
    assert len([r for r in view.rows if r.exposure == 'E3']) == 2
    assert sum('待审试算' in row.remark for row in view.rows) == 1
    assert state == before


def test_incomplete_branch_tuple_never_combines_known_dimensions(active_method):
    state = mixed_pending_group()
    state.scenarios = state.scenarios[:1]
    state.risk_results = state.risk_results[:1]
    state.risk_results[0].branch_evaluations = [
        {'branch_id': 'first', 'S': 'S2', 'E': None, 'C': 'C2', 'ASIL': None},
        {'branch_id': 'second', 'S': None, 'E': 'E3', 'C': 'C3', 'ASIL': None}]
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    row = view.rows[0]
    assert (row.severity, row.exposure, row.controllability, row.asil) == ('S2', '待评定', 'C2', '待评定')
    assert view.projection_metrics['grouping_map'][0]['member_ratings'][0]['branches'] == state.risk_results[0].branch_evaluations
    assert state == before


@pytest.mark.parametrize('input_key', ['override_inputs', 'physical_inputs'])
def test_outside_driver_profile_gap_is_explained_as_inapplicability_not_missing_ttc(input_key):
    text = EngineeringReportTextMapper.controllability_branch({'controllability': {
        'decision_status': 'PROFILE_APPLICABILITY_UNRESOLVED',
        'pending_reason': 'TTC_PROFILE_APPLICABILITY_UNRESOLVED: driver_in_vehicle',
        input_key: {'driver_in_vehicle': False}, 'derived_ttc': {'ttc_s': 5.9}}})
    assert '驾驶员在车外' in text and '不适用' in text
    assert '缺相对距离' not in text and 'C0' not in text


def test_pending_asil_explains_same_branch_matrix_without_assigning_c(active_method):
    state = mixed_pending_group()
    risk = state.risk_results[0]
    risk.severity = EvidenceValue('S3', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    risk.exposure = EvidenceValue('E3', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    detail = next(d for d in view.scenario_details if d.scenario_id == risk.scenario_id)
    assert detail.asil == detail.controllability == '待评定'
    assert 'S3/E3' in detail.pending_reasons and 'C3→C' in detail.pending_reasons
    assert '未选择C' in detail.pending_reasons
    assert '14条风险明细组织为3个审核组' in view.projection_metrics['grouping_explanation']
    assert state == before


def test_existing_safe_state_draft_survives_projection_and_reaggregation(active_method):
    from hara_agent.services.analysis import MethodSafetyGoalService
    from hara_agent.workflow.nodes.safety_goals import aggregate_safety_goals
    state = mixed_pending_group()
    aggregate_safety_goals(state, MethodSafetyGoalService(active_method))
    goal = state.safety_goals[0]
    goal.safe_state = '车辆停止并保持静止'
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert view.safety_goals[0].safe_state == '候选：车辆停止并保持静止'
    assert goal.status is ReviewStatus.PENDING and state == before
    aggregate_safety_goals(state, MethodSafetyGoalService(active_method))
    assert state.safety_goals[0].safe_state == goal.safe_state
    assert state.safety_goals[0].status is ReviewStatus.PENDING


def test_incomplete_group_shared_value_and_branch_difference(active_method):
    state = mixed_pending_group()
    for risk in state.risk_results:
        risk.exposure = EvidenceValue('E3', ReviewStatus.FINALIZED, rule_version='fixture')
    project = HARAReportProjectionService(load_report_schema())
    assert project.project(state, active_method).rows[0].exposure == 'E3'
    state.risk_results[0].branch_evaluations = [{'S':'S1', 'E':'E2', 'C':None, 'ASIL':None}]
    view = project.project(state, active_method)
    assert {r.exposure for r in view.rows} == {'E2', 'E3'}
    assert view.scenario_details[0].exposure == 'E2'
    assert state.risk_results[0].exposure.value == 'E3'  # Stored aggregate is untouched.


def test_incomplete_group_keeps_different_actual_exposure_bases(active_method):
    state = mixed_pending_group()
    trace = {'assessments': []}
    for i, risk in enumerate(state.risk_results):
        record = {'malfunction_id': risk.malfunction_id, 'scenario_id': risk.scenario_id,
                  'risk_scoring_invoked': True}
        for field in ('severity', 'exposure', 'controllability', 'asil'):
            evidence = getattr(risk, field)
            record[field] = {'result': evidence.value, 'status': evidence.status.value}
        record['exposure'].update(
            atom_bindings=[{'atom_id': f'place-{i}', 'used': True, 'e_rank': risk.exposure.value, 'dimensions': ['WHERE']}],
            scenario_terms={'WHERE': {'atom_id': f'place-{i}',
                'atom_label': '室外停车场' if i < 12 else '地下停车场', 'project_value': '来源允许的停车场'}},
            aggregation_rule='min_when_unequal')
        trace['assessments'].append(record)
    before = deepcopy((state, trace))
    rows = HARAReportProjectionService(load_report_schema()).project(state, active_method, risk_trace=trace).rows
    assert len(rows) == 2
    by_e = {r.exposure: r for r in rows}
    assert '室外停车场=E3' in by_e['E3'].exposure_rationale
    assert '地下停车场=E2' in by_e['E2'].exposure_rationale
    assert all('不同暴露等级组合取较低等级' in r.exposure_rationale for r in rows)
    assert (state, trace) == before


@pytest.mark.parametrize('dimension,grades', [
    ('severity', ('S1', 'S2')), ('exposure', ('E2', 'E3')),
    ('controllability', ('C1', 'C2')),
])
def test_group_dimension_distinguishes_mixed_values_from_missing(active_method, dimension, grades):
    state = mixed_pending_group()
    for i, risk in enumerate(state.risk_results):
        setattr(risk, dimension, EvidenceValue(grades[i % 2], ReviewStatus.FINALIZED, rule_version='fixture'))
    project = HARAReportProjectionService(load_report_schema())
    before = deepcopy(state)
    rows = project.project(state, active_method).rows
    assert {getattr(row, dimension) for row in rows} == set(grades)
    assert all(row.asil == '待评定' for row in rows) and state == before
    setattr(state.risk_results[0], dimension, _pending())
    before = deepcopy(state)
    view = project.project(state, active_method)
    assert {getattr(row, dimension) for row in view.rows} == {*grades, '待评定'}
    assert sum(getattr(d, dimension) == '待评定' for d in view.scenario_details) == 1
    assert state == before


def test_group_unknown_collision_branch_does_not_hide_known_e(active_method):
    state = mixed_pending_group()
    state.risk_results[0].branch_evaluations = [{'S':'S1', 'E':None, 'C':None, 'ASIL':None}]
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert {row.exposure for row in view.rows} == {'E2', 'E3', '待评定'}
    assert '分支评级尚未闭合' in view.rows[0].remark
    assert state == before


def test_tied_highest_asil_keeps_one_actual_tuple_and_all_other_tuples(active_method):
    state = mixed_pending_group()
    state.risk_results = state.risk_results[:2]
    state.scenarios = state.scenarios[:2]
    tuples = [('S1', 'E3', 'C3', 'A'), ('S2', 'E2', 'C3', 'A')]
    for risk, values in zip(state.risk_results, tuples):
        for field, value in zip(('severity', 'exposure', 'controllability', 'asil'), values):
            setattr(risk, field, EvidenceValue(value, ReviewStatus.FINALIZED, rule_version='fixture'))
        risk.branch_evaluations = [dict(zip(('S', 'E', 'C', 'ASIL'), values))]
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert len(view.rows) == 2
    assert {(r.severity, r.exposure, r.controllability, r.asil) for r in view.rows} == set(tuples)
    assert {(d.severity, d.exposure, d.controllability, d.asil) for d in view.scenario_details} == set(tuples)
    assert [m['branches'] for g in view.projection_metrics['grouping_map'] for m in g['member_ratings']] == [r.branch_evaluations for r in state.risk_results]
    assert state == before


@pytest.mark.parametrize('value,conditions', [
    ('C1', {'remote_intervention_available': True}),
    ('C3', {'driver_in_vehicle': False, 'remote_intervention_available': False,
            'other_road_user_avoidance_possible': False}),
])
def test_c_rationale_uses_matched_override_even_when_ttc_exists(value, conditions):
    evidence = EvidenceValue(value, ReviewStatus.FINALIZED, rule_version='fixture')
    trace = {'decision_status':'OVERRIDE_MATCH', 'override_inputs':conditions,
             'derived_ttc':{'ttc_s':6}}
    before = deepcopy((evidence, trace))
    text = EngineeringReportTextMapper().controllability_rationale(evidence,trace)
    assert f'优先干预规则评为 {value}' in text
    assert '未使用 TTC 分档' in text and '按当前纵向' not in text
    assert (evidence, trace) == before


def test_c_ttc_rationale_keeps_outside_driver_applicability_visible():
    evidence = EvidenceValue('C0',ReviewStatus.PENDING,review_reason='CONDITIONAL_TRIAL')
    trace = {'decision_status':'TTC_AFTER_UNKNOWN_OVERRIDE', 'unknown_policy_action':'SKIP_TO_TTC',
             'override_inputs':{'driver_in_vehicle':False,'remote_intervention_available':None,
                                'other_road_user_avoidance_possible':None},
             'derived_ttc':{'ttc_s':5.9115}}
    before = deepcopy((evidence,trace))
    text = EngineeringReportTextMapper().controllability_rationale(evidence,trace)
    assert 'C0' in text and '5.91' in text
    assert '驾驶员在车外' in text and '未证明实际可用的干预路径' in text
    assert '无远程干预' not in text
    assert (evidence,trace) == before


def test_c_rationale_uses_branch_driver_context_instead_of_parent():
    parent = {'controllability':{'override_inputs':{'driver_in_vehicle':True}}}
    branch = {'S':'S1','E':'E3','C':'C0','ASIL':'QM','ttc_s':5.9115,
              'risk_context':{'driver_in_vehicle':{'status':'AVAILABLE','value':False}},
              'scores':{'controllability':{'decision_status':'TTC_AFTER_UNKNOWN_OVERRIDE',
                        'unknown_policy_action':'SKIP_TO_TTC'}}}
    before = deepcopy((parent,branch))
    displayed = HARAReportProjectionService._display_trace('controllability',parent,branch)
    text = EngineeringReportTextMapper().controllability_rationale(
        EvidenceValue('C0',ReviewStatus.PENDING,review_reason='CONDITIONAL_TRIAL'),displayed)
    assert '驾驶员在车外' in text and '未证明实际可用的干预路径' in text
    assert 'C0' in text and (parent,branch)==before


def test_concrete_severity_gap_is_not_hidden_by_input_production_status():
    mapper = EngineeringReportTextMapper()
    trace = {'pending_reason': 'MISSING_RELATIVE_SPEED',
             'input_production': {'pending_reason': 'AUTONOMOUS_MODEL_UNRESOLVED', 'after': {}}}
    before = deepcopy(trace)
    text = mapper.severity_rationale(_pending(), trace)
    assert '碰撞相对速度' in text and '输入生产尚有未决项' not in text
    assert trace == before
    trace['input_production']['pending_reason'] = 'MOTION_MODEL_OR_ENVELOPE_NOT_SOURCED'
    text = mapper.severity_rationale(_pending(), trace)
    assert '碰撞相对速度' in text and '适用运动模型或速度包络' in text


def test_exposure_rank_labels_the_used_atom_not_the_search_phrase():
    mapper = EngineeringReportTextMapper(load_report_schema().presentation['language'])
    trace = {'result': 'E3', 'aggregation_rule': 'e3_e4_mix',
             'atom_bindings': [{'atom_id': 'FB006', 'used': True, 'E_class': 'E3', 'dimension': ['ROAD']}],
             'scenario_terms': {'ROAD': {'atom_id': 'FB006', 'atom_label': 'Reduced friction (µ < 0.8)',
                                        'project_value': '晴朗'}}}
    before = deepcopy(trace)
    evidence = EvidenceValue('E3', ReviewStatus.FINALIZED, rule_version='fixture')
    text = mapper.exposure_rationale(evidence, trace)
    assert '低附着路面（µ < 0.8）=E3' in text and '晴朗' not in text
    assert trace == before
    # A stale label for a different Atom must not explain this Atom's rank.
    trace['scenario_terms']['ROAD']['atom_id'] = 'different-atom'
    text = mapper.exposure_rationale(evidence, trace)
    assert '道路条件=E3' in text and '低附着' not in text


def test_full_function_inventory_includes_functions_without_risk_rows(active_method):
    state = mixed_pending_group()
    state.functions.append({'function_id': 'excluded-function', 'name': '仅排除的功能', 'output': '信号', 'status': 'PENDING'})
    state.malfunctions.append({'function_id': 'excluded-function', 'malfunction_id': 'excluded-mf', 'description': '失效'})
    trace = {'assessments': [dict(malfunction_id='excluded-mf', scenario_id='excluded-scene',
                                risk_eligibility={'status': 'INELIGIBLE_CAUSAL'}, risk_scoring_invoked=False)]}
    # This fixture has no scored trace: derive the complete upstream inventory
    # independently of the rendered HARA rows, then check exclusion accounting.
    state.risk_results = []
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method, risk_trace=trace)
    assert len(view.function_inventory) == len(state.functions)
    function = next(f for f in view.function_inventory if f.function_id == 'excluded-function')
    assert (function.malfunctions, function.analyzed, function.retained, function.excluded, function.pending) == (1, 1, 0, 1, 0)
    assert function.status == 'PENDING' and not view.rows
    assert state == before


def test_malfunction_distribution_keeps_numeric_pending_and_known_severity(active_method):
    state = mixed_pending_group()
    state.risk_results[0].severity = EvidenceValue('S3', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    state.risk_results[0].exposure = EvidenceValue('E3', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    inventory = next(m for m in view.malfunction_inventory if m.malfunction_id == state.risk_results[0].malfunction_id)
    assert '已知最大 S=S3' in inventory.risk_distribution
    assert 'S3×1' in inventory.risk_distribution and '待审试算×1' in inventory.risk_distribution
    assert 'E3×12' in inventory.risk_distribution
    assert view.rows[0].severity == 'S3'
    assert state == before


def test_summary_separates_review_groups_and_retained_records(tmp_path, active_method):
    schema = load_report_schema()
    view = HARAReportProjectionService(schema).project(mixed_pending_group(), active_method)
    output = tmp_path / 'report.xlsx'
    HARAReportWorkbookRenderer().render(view, 'references/HARA_Template_AI_20260327.xlsx', output, schema)
    workbook = load_workbook(output)
    summary = {row[1].value: row[2].value for row in workbook['00_Summary'].iter_rows(min_row=5)}
    assert summary['主表审核组数（04_HARA）'] == 2
    assert summary['风险记录数'] == summary['场景明细成员数'] == 14
    assert view.summary.eligible_hazardous_event_count == 14


def test_sg_rewording_requires_all_member_sources_and_keeps_status(active_method):
    schema = load_report_schema()
    mapper = EngineeringReportTextMapper(schema.presentation['language'])
    rules = schema.presentation['language']['safety_goal_rewordings']
    for rule in rules:
        member = {'scenario_id': 'A', 'malfunction_id': 'MF', **rule['source']}
        result = mapper.safety_goal_presentation([member], '原目标')
        assert result['text'] == rule['text']
        assert result['members'][0]['basis'] == rule['source']
        changed = {**member, 'scenario_id': 'B', 'vehicle_hazard': '另一种危险机制', 'malfunction': '另一种失效'}
        result = mapper.safety_goal_presentation([member, changed], '原目标')
        assert '原目标' in result['text']  # unsupported wording retains its original goal
        assert result['members'][1]['malfunction'] == '另一种失效'
        assert result['members'][1]['wording_id'] == 'grounded_behavior_excerpt'
    # A claim about absent contact must not get rewritten into a collision goal.
    changed = {**rules[1]['source'], 'hazardous_event': '未发生接触'}
    displayed = mapper.safety_goal_presentation([changed], '原目标')
    assert displayed['members'][0]['wording_id'] == 'grounded_behavior_excerpt'
    assert '碰撞' not in displayed['text']
    state = mixed_pending_group()
    state.malfunctions[0].update(description=rules[0]['source']['malfunction'],
        vehicle_level_hazard=rules[0]['source']['vehicle_hazard'])
    state.risk_results[0].safety_goal_id = 'SG-1'
    state.safety_goals = [SafetyGoal('SG-1', '原目标', '待确定', 'B', ['SCN-0'], ReviewStatus.PENDING)]
    before = deepcopy(state)
    view = HARAReportProjectionService(schema).project(state, active_method)
    assert view.safety_goals[0].safety_goal == rules[0]['text']
    assert view.safety_goals[0].original_text == '原目标'
    assert state == before


def test_harm_translation_preserves_qualifiers_without_inferring_missing_harm():
    mapper = EngineeringReportTextMapper(load_report_schema().presentation['language'])
    text = mapper.potential_harm('Severe injuries, possibly life-threatening, survival probable', severity=_pending())
    assert '可能危及生命' in text and '生存可能性较高' in text
    assert mapper.potential_harm('unknown source phrase', severity=_pending()) == 'unknown source phrase'
    assert mapper.potential_harm('', severity=_pending()) == '待S评定完成后确定'


@pytest.mark.parametrize('known,pending', [(4, 4), (4, 20), (8, 14)])
def test_qm_with_unresolved_members_is_not_a_group_qm(active_method, known, pending):
    state = mixed_pending_group()
    source, assessment = state.scenarios[0], state.risk_results[0]
    causal = state.item_definition['scenario_assessments'][0]
    state.item_definition['scenario_assessments'] = []
    state.scenarios, state.risk_results = [], []
    for i in range(known + pending):
        child, risk = deepcopy(source), deepcopy(assessment)
        child.scenario_id = risk.scenario_id = f'SCN-{i}'
        risk.assessment_id = f'RA-{i}'
        if i < known:
            risk.severity = EvidenceValue('S0', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
            risk.asil = EvidenceValue('QM', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
        state.scenarios.append(child)
        state.risk_results.append(risk)
        state.item_definition['scenario_assessments'].append({**deepcopy(causal), 'scenario_id': child.scenario_id})
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert len(view.rows) == 2
    assert {r.asil for r in view.rows} == {'QM', '待评定'}
    for row in view.rows:
        members = [d for d in view.scenario_details if d.hara_id == row.hara_id]
        assert {d.asil for d in members} == {row.asil}
        assert len(members) == (known if row.asil == 'QM' else pending)
    assert Counter(d.asil for d in view.scenario_details) == {'QM': known, '待评定': pending}
    assert state == before


@pytest.mark.parametrize('known_grade,pending', [('QM', False), ('A', True)])
def test_group_qm_policy_keeps_known_matrix_results(active_method, known_grade, pending):
    state = mixed_pending_group()
    for i, risk in enumerate(state.risk_results):
        if not pending or i == 0:
            risk.asil = EvidenceValue(known_grade, ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert view.rows[0].asil == known_grade  # C missing alone does not invalidate known QM.
    assert state == before


def test_qm_group_keeps_unresolved_collision_branch_visible(active_method):
    state = mixed_pending_group()
    for risk in state.risk_results:
        risk.asil = EvidenceValue('QM', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    state.risk_results[0].branch_evaluations = [
        {'S': 'S0', 'E': 'E3', 'C': None, 'ASIL': 'QM'},
        {'S': 'S1', 'E': 'E3', 'C': None, 'ASIL': None},
    ]
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    assert view.rows[0].asil == '待评定'
    assert '碰撞分支' in view.rows[0].asil_rationale
    assert state == before


@pytest.mark.parametrize('goal_ids', [('SG-A',), ('SG-A', 'SG-B')])
def test_group_sg_links_include_nonrepresentatives_without_creating_qm_goals(active_method, tmp_path, goal_ids):
    state = mixed_pending_group()
    state.risk_results[0].asil = EvidenceValue('QM', ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL')
    state.safety_goals = []
    for i, sg_id in enumerate(goal_ids, start=1):
        risk = state.risk_results[i]
        risk.safety_goal_id = sg_id
        state.safety_goals.append(SafetyGoal(sg_id, f'原目标{i}', '待确定', '', [risk.scenario_id], ReviewStatus.PENDING))
    # Repeated references identify one existing goal, never another goal.
    state.risk_results[-1].safety_goal_id = goal_ids[0]
    before = deepcopy(state)
    schema = load_report_schema()
    view = HARAReportProjectionService(schema).project(state, active_method)
    assert not view.rows[0].sg_id  # QM member keeps its original no-SG behavior.
    row = next(r for r in view.rows if r.sg_id.splitlines() == list(goal_ids))
    assert row.sg_id.splitlines() == list(goal_ids)
    for goal in view.safety_goals:
        assert goal.safety_goal in row.safety_goal
        if len(goal_ids) > 1:
            assert f'{goal.sg_id}：' in row.safety_goal and f'{goal.sg_id}：' in row.safe_state
    members = [m for g in view.projection_metrics['grouping_map'] for m in g['member_ratings']]
    assert {m['scenario_id']: m['safety_goal_id'] for m in members} == {
        risk.scenario_id: risk.safety_goal_id for risk in state.risk_results}
    output = tmp_path / 'group_sg.xlsx'
    HARAReportWorkbookRenderer().render(view, 'references/HARA_Template_AI_20260327.xlsx', output, schema)
    workbook = load_workbook(output)
    main_text = '\n'.join(str(cell.value) for cells in workbook['04_HARA'] for cell in cells if cell.value)
    for sg_id in goal_ids:
        assert sg_id in main_text
    assert not state.risk_results[0].safety_goal_id
    assert state == before


def test_ftti_extrema_describe_discrete_contact_samples_only():
    timings = [{'branch_timings': [
        {'contact_time_s': 0.154286, 'contact_time_kind': 'ttc_s', 'time_origin': 'MODEL_INITIAL_STATE'},
        {'contact_time_s': 3.0, 'contact_time_kind': 'ttc_s', 'time_origin': 'MODEL_INITIAL_STATE'},
        {'contact_time_s': None},
    ]}]
    before = deepcopy(timings)
    text = EngineeringReportTextMapper.ftti_rationale(timings)
    assert '最小 0.15 s' in text and '最大 3 s' in text
    assert '已计算样本' in text and '连续' in text and '不能作为 FTTI' in text
    assert '[' not in text and '端点包含性' not in text
    assert timings == before


@pytest.mark.parametrize('decision,expected', [
    ('MODEL_NOT_APPLICABLE', '当前方法无适用 C 规则'),
    ('MODEL_DEPENDENCY_UNRESOLVED', 'C 方法依赖尚未成立'),
])
def test_control_method_gap_is_not_reported_as_missing_geometry(decision, expected):
    trace = {'decision_status': decision, 'input_diagnostic': {'reason': 'MISSING_DISTANCE'},
             'known_physical_inputs': {'ttc_s': 1.2345}}
    before = deepcopy(trace)
    text = EngineeringReportTextMapper().controllability_rationale(_pending(), trace)
    assert expected in text and '缺相对距离' not in text
    assert '1.23' in text and trace == before


def test_conditional_numeric_results_are_not_missing_inputs_in_inventory(active_method):
    state = mixed_pending_group()
    for risk in state.risk_results:
        for field, value in (('severity', 'S1'), ('exposure', 'E3'), ('controllability', 'C3'), ('asil', 'A')):
            setattr(risk, field, EvidenceValue(value, ReviewStatus.PENDING, review_reason='CONDITIONAL_TRIAL'))
    before = deepcopy(state)
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method)
    inventory = view.malfunction_inventory[0]
    assert inventory.pending_summary == ''
    assert '条件性试算' in inventory.risk_distribution
    assert state == before


def test_inventory_uses_missing_dimension_method_reason(active_method):
    state = mixed_pending_group()
    trace = {'assessments': []}
    for risk in state.risk_results:
        row = dict(malfunction_id=risk.malfunction_id, scenario_id=risk.scenario_id, risk_scoring_invoked=True)
        for field in ('severity', 'exposure', 'controllability', 'asil'):
            evidence = getattr(risk, field)
            row[field] = dict(result=evidence.value, status=evidence.status.value)
        row['controllability'].update(decision_status='MODEL_NOT_APPLICABLE', input_diagnostic={'reason': 'TTC_AVAILABLE'})
        trace['assessments'].append(row)
    before = deepcopy((state, trace))
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method, risk_trace=trace)
    assert '当前方法无适用 C 规则' in view.malfunction_inventory[0].pending_summary
    assert 'TTC_AVAILABLE' not in view.malfunction_inventory[0].pending_summary
    assert (state, trace) == before


def test_excluded_function_keeps_actual_causal_reason_without_missing_risk_claim(active_method):
    state = mixed_pending_group()
    risk = state.risk_results[0]
    state.risk_results = []
    rationale = '缺少证据表明该额外开启的功能会触发纵向加速。'
    state.item_definition['scenario_assessments'][0]['rationale'] = rationale
    trace = {'assessments': [dict(malfunction_id=risk.malfunction_id, scenario_id=risk.scenario_id,
        risk_scoring_invoked=False, risk_eligibility={'status': 'INELIGIBLE_CAUSAL', 'reason_codes': ['CAUSAL_GAP']})]}
    before = deepcopy((state, trace))
    view = HARAReportProjectionService(load_report_schema()).project(state, active_method, risk_trace=trace)
    assert view.generated_children[0].disposition_reason == rationale
    inventory = view.malfunction_inventory[0]
    assert inventory.disposition == 'EXCLUDED_FROM_RISK' and inventory.excluded == 1
    assert '未进入风险评分' in inventory.risk_distribution
    assert '待评定' not in inventory.risk_distribution and 'Pending' not in inventory.risk_distribution
    assert (state, trace) == before
