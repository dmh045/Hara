from copy import deepcopy
from dataclasses import replace
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from hara_agent.models import (
    EvidenceValue, ReviewStatus, RiskAssessment, ScenarioCandidate, SourceRef,
)
from hara_agent.services.reporting import (
    EngineeringReportTextMapper, HARAReportProjectionService,
    audit_content_presentation, load_report_schema,
    load_scenario_projection_contexts,
)
from hara_agent.workflow.state import HARAState


ROOT = Path(__file__).resolve().parents[2]


def _pending() -> EvidenceValue:
    return EvidenceValue(None, ReviewStatus.PENDING, review_reason="upstream pending")


def _scenario(index: int, coverage: str = "typical") -> ScenarioCandidate:
    return ScenarioCandidate(
        scenario_id=f"SCN-{index}",
        operating_scenario="室内停车场",
        situational_description="",
        situational_detailing="",
        operating_mode="active",
        facts={"ego_speed_constraint": {"min_kph": 0.0, "max_kph": 5.0}},
        sources=[SourceRef("item_definition", "ItemDef.docx", f"paragraph[{index}]")],
        source_scenario_id="SCN-PARENT",
        atomic_variant=f"scenario_synthesis:{coverage}",
        context_resolution={"scenario_synthesis": {"coverage_label": coverage}},
        analysis_instance={
            "semantic_group_id": "SYNTH-GROUP-1",
            "parent_scenario_id": "SCN-PARENT",
            "hazardous_event_id": "HE-1",
            "selected_atoms": [f"FA00{index}"],
            "structured_semantic_query": {
                "location_categories": ["LOCATION_PARKING"],
                "action_categories": ["ACTION_PARK"],
                "object_categories": ["OBJECT_PEDESTRIAN"],
            },
            "validation_status": "VALIDATED",
        },
    )


def _risk(index: int, hazardous_event: str = "车辆可能与行人碰撞") -> RiskAssessment:
    return RiskAssessment(
        assessment_id=f"RA-{index}", scenario_id=f"SCN-{index}",
        severity=_pending(), exposure=_pending(), controllability=_pending(),
        asil=_pending(), malfunction_id="MF-1",
        hazardous_event=hazardous_event,
    )


def _state(hazardous_events: tuple[str, ...]) -> HARAState:
    coverage = ("typical", "boundary", "extreme")
    return HARAState(
        run_id="reviewer-report",
        functions=[{"function_id": "F-1", "name": "自主泊车", "output": "车辆运动"}],
        malfunctions=[{
            "malfunction_id": "MF-1", "function_id": "F-1", "guideword": "丧失",
            "description": "行人识别丧失", "vehicle_level_hazard": "车辆未及时制动",
        }],
        scenarios=[_scenario(index, coverage[index - 1]) for index in range(1, 4)],
        risk_results=[
            _risk(index, hazardous_events[index - 1]) for index in range(1, 4)
        ],
    )


def _method() -> SimpleNamespace:
    return SimpleNamespace(
        metadata={"method_source_hash": "method-hash"},
        guidewords=SimpleNamespace(guidewords=["丧失"]),
    )


def test_three_siblings_project_to_one_main_row_and_three_detail_rows_without_state_mutation():
    state = _state(("车辆可能与行人碰撞",) * 3)
    before = deepcopy(state)

    view = HARAReportProjectionService(load_report_schema()).project(state, _method())

    assert len(view.rows) == 1
    assert len(view.scenario_details) == 3
    assert {item.hara_id for item in view.scenario_details} == {view.rows[0].hara_id}
    assert {item.semantic_group_id for item in view.scenario_details} == {"SYNTH-GROUP-1"}
    assert all("行人" in item.object_interaction_summary for item in view.scenario_details)
    assert len(view.audit_references) == 3
    assert all(not item.risk_trace_reference for item in view.audit_references)
    assert view.projection_metrics["grouping_reduction"] == 2
    assert state == before


def test_hazardous_event_text_divergence_prevents_false_merge():
    state = _state((
        "车辆可能与行人碰撞",
        "车辆可能与行人碰撞",
        "车辆可能撞击静态障碍物",
    ))

    view = HARAReportProjectionService(load_report_schema()).project(state, _method())

    assert len(view.rows) == 2
    assert len(view.scenario_details) == 3
    assert view.projection_metrics[
        "groups_not_merged_due_hazardous_event_divergence"
    ] == 1


def test_text_mapper_preserves_range_and_uses_only_finalized_point_speed():
    mapper = EngineeringReportTextMapper()
    scenario = _scenario(1)
    scenario.facts["ego_speed_constraint"] = {"min_kph": 0.0, "max_kph": 7.0}
    assert mapper.speed_text(scenario) == "适用车速范围：0–7 km/h"

    scenario.facts["ego_speed_kph"] = 3.5
    scenario.fact_provenance["ego_speed_kph"] = {"approval": "FINALIZED"}
    assert mapper.speed_text(scenario) == "分析车速：3.5 km/h"


def test_hazardous_event_compactor_only_normalizes_source_terms_and_duplicates():
    source = "AVP在Active模式；AVP在Active模式；车辆未拉起EPB"

    rendered = EngineeringReportTextMapper().hazardous_event(source)

    assert rendered == "AVP在激活模式；车辆未拉起电子驻车制动。"
    assert "车辆未拉起" in rendered


def test_content_gate_rejects_raw_dimension_atom_and_unapproved_language():
    state = _state(("车辆可能与行人碰撞",) * 3)
    view = HARAReportProjectionService(load_report_schema()).project(state, _method())
    bad_row = replace(
        view.rows[0],
        operational_scenario="WHERE=Garage；FA001 | Drive low speed。",
    )
    bad_view = replace(view, rows=(bad_row,))

    audit = audit_content_presentation(
        bad_view, {"classification": "UPSTREAM_RISK_NOT_READY"},
    )

    assert audit["quality_gate"] == "FAIL"
    assert audit["raw_dimension_syntax_leakage_count"] == 1
    assert audit["raw_atom_id_leakage_count"] == 1
    assert audit["language_mix_count"] == 1


def test_content_gate_allows_epb_as_an_engineering_acronym():
    state = _state(("车辆未拉起EPB并可能发生溜车",) * 3)
    view = HARAReportProjectionService(load_report_schema()).project(state, _method())
    row = replace(view.rows[0], hazardous_event="车辆未拉起EPB并可能发生溜车。")

    audit = audit_content_presentation(
        replace(view, rows=(row,)),
        {"classification": "UPSTREAM_RISK_NOT_READY"},
    )

    assert audit["language_mix_count"] == 0


def test_fresh_report_context_restores_group_coverage_and_contextual_speed(tmp_path):
    state = _state(("车辆可能与行人碰撞",) * 3)
    state.scenarios = state.scenarios[:1]
    state.risk_results = state.risk_results[:1]
    before = deepcopy(state)
    key = ("MF-1", "SCN-1", "HE-1")
    candidates = tmp_path / "scenario_synthesis_candidates.json"
    candidates.write_text(json.dumps({"groups": [{
        "malfunction_id": key[0],
        "parent_scenario_id": key[1],
        "hazardous_event_id": key[2],
        "semantic_group_id": "SYNTH-FRESH-1",
        "structured_semantic_query": {
            "location_categories": ["LOCATION_PARKING"],
            "action_categories": ["ACTION_PARK"],
            "object_categories": ["OBJECT_PEDESTRIAN"],
            "traffic_relations": ["TRAFFIC_CROSSING"],
            "road_relations": [],
        },
        "coverage_plan": {
            "desired_variant_count": 3,
            "variant_intents": [
                {"coverage_label": "typical"},
                {"coverage_label": "boundary"},
                {"coverage_label": "extreme"},
            ],
        },
    }]}), encoding="utf-8")
    speed = tmp_path / "speed_context.json"
    speed.write_text(json.dumps({"records": [{
        "malfunction_id": key[0],
        "parent_scenario_id": key[1],
        "hazardous_event_id": key[2],
        "classification": "CONTEXTUAL_SPEED_CONSUMED",
        "selected_context": "PARKING",
        "match_basis": "FUNCTION.odd_constraints[0]",
        "report_visible_range": [0.0, 5.0],
        "source_speed_envelopes": [],
    }]}), encoding="utf-8")
    contexts = load_scenario_projection_contexts(
        synthesis_candidates_path=candidates,
        speed_context_audit_path=speed,
    )

    view = HARAReportProjectionService(load_report_schema()).project(
        state,
        _method(),
        risk_trace={"assessments": [{
            "malfunction_id": key[0],
            "scenario_id": key[1],
            "hazardous_event_id": key[2],
        }]},
        scenario_projection_contexts=contexts,
    )

    detail = view.scenario_details[0]
    assert detail.semantic_group_id == "SYNTH-FRESH-1"
    assert detail.variant == "计划覆盖：代表场景 / 边界场景 / 高要求场景"
    assert detail.speed_constraint == "适用车速范围：不高于5 km/h"
    assert "车辆执行泊车" in detail.operational_scenario
    assert "横穿交通参与者" in detail.operational_scenario
    assert "当前为综合前审阅投影" in view.rows[0].scenario_detail
    assert view.projection_metrics["contextual_speed_rows"] == 1
    assert state == before


def test_driver_branches_and_trial_values_remain_distinct_in_formal_report(tmp_path):
    from hara_agent.services.reporting import HARAReportWorkbookRenderer

    def finalized(value: str) -> EvidenceValue:
        return EvidenceValue(value, ReviewStatus.FINALIZED, rule_version="method-v1")

    def risk(scenario_id: str, *, trial: bool) -> RiskAssessment:
        return RiskAssessment(
            assessment_id=f"RA-{scenario_id}", scenario_id=scenario_id,
            severity=finalized("S2"),
            exposure=(
                EvidenceValue("E3", ReviewStatus.PENDING, review_reason="CANDIDATE_DEFAULT")
                if trial else finalized("E3")
            ),
            controllability=(
                EvidenceValue("C2", ReviewStatus.PENDING, review_reason="CONDITIONAL_TRIAL")
                if trial else finalized("C2")
            ),
            asil=(
                EvidenceValue("B", ReviewStatus.PENDING, review_reason="ASIL depends on S/E/C")
                if trial else finalized("B")
            ),
            malfunction_id="MF-1", hazardous_event="车辆可能与行人碰撞",
        )

    def scenario(position: str) -> ScenarioCandidate:
        item = _scenario(1)
        item.scenario_id = f"SCN-DRIVER-{position}"
        item.facts.update({
            "driver_position": position,
            "ego_speed_kph": 7,
            "object_speed_kph": 0,
            "relative_speed_kph": 7,
            "relative_distance_m": 3,
            "object_position": "front",
            "ego_longitudinal_direction": "FORWARD",
            "object_longitudinal_direction": "STATIONARY",
        })
        item.fact_provenance["ego_speed_kph"] = {
            "approval": "FINALIZED", "policy_id": "AVP_EGO_SPEED_CLOSED_UPPER_BOUND_V1",
            "source_refs": [{"location": "table[14].row[15]"}],
        }
        item.analysis_instance["driver_configuration_branch"] = {
            "driver_position": position,
        }
        return item

    state = _state(("车辆可能与行人碰撞",) * 3)
    state.scenarios = [scenario("in_driver_seat"), scenario("outside_driver_seat")]
    state.risk_results = [
        risk("SCN-DRIVER-in_driver_seat", trial=False),
        risk("SCN-DRIVER-outside_driver_seat", trial=True),
    ]
    def trace(scenario_id: str, *, trial: bool) -> dict:
        return {
            "malfunction_id": "MF-1", "scenario_id": scenario_id,
            "risk_scoring_invoked": True,
            "severity": {"status": "FINALIZED", "result": "S2"},
            "exposure": {"status": "PENDING_INPUT" if trial else "FINALIZED", "result": "E3"},
            "asil": {"status": "PENDING_UPSTREAM_RISK_VALUE" if trial else "FINALIZED", "result": "B"},
            "hazardous_event_risk_context": {
                "ego_speed_kph": {"status": "AVAILABLE", "value": 7},
                "object_speed_kph": {"status": "AVAILABLE", "value": 0},
                "relative_speed_kph": {"status": "AVAILABLE", "value": 7},
                "relative_distance_m": {"status": "AVAILABLE", "value": 3},
                "ttc_s": {"status": "AVAILABLE", "value": 1.54},
                "driver_in_vehicle": {
                    "status": "AVAILABLE" if "in_driver" in scenario_id else "PENDING",
                    "value": True if "in_driver" in scenario_id else None,
                },
            },
            "controllability": {
                "status": "PENDING_INPUT" if trial else "FINALIZED", "result": "C2",
                "decision_tree_stage": "TTC", "rule_ids": ["C-TTC-01"],
                "unknown_override_policy": "UNSPECIFIED",
                "derived_ttc": {"closing_speed_kph": 7, "ttc_s": 1.54},
            },
        }

    schema = load_report_schema()
    view = HARAReportProjectionService(schema).project(
        state, _method(), risk_trace={"assessments": [
            trace("SCN-DRIVER-in_driver_seat", trial=False),
            trace("SCN-DRIVER-outside_driver_seat", trial=True),
        ]},
    )
    assert len(view.rows) == 2
    assert "驾驶员在驾驶位" in view.rows[0].operational_scenario
    assert "驾驶员不在驾驶位" in view.rows[1].operational_scenario
    assert view.rows[0].remark == "S 已计算；E 已计算；C 已计算；ASIL 已计算"
    assert view.rows[1].remark == "S 已计算；E 待审试算；C 待审试算；ASIL 待审试算"
    assert view.rows[1].exposure == "E3（试算）"
    assert view.rows[1].asil == "B（试算）"
    assert "待审分析设定" in view.rows[1].exposure_rationale
    assert view.projection_metrics["score_status_counts"]["asil"] == {
        "calculated": 1, "conditional_trial": 1, "uncalculated": 0,
    }
    assert "自车速度：7 km/h" in view.scenario_details[0].physical_inputs
    assert "TTC：1.54 s" in view.scenario_details[0].physical_inputs
    assert "目标位置：前方" in view.scenario_details[0].physical_inputs
    assert "车内状态：未确定" in view.scenario_details[1].driver_branch
    assert "命中规则：C-TTC-01" in view.scenario_details[0].controllability_branch
    assert "table[14].row[15]" in view.scenario_details[0].analysis_basis

    output = tmp_path / "review.xlsx"
    HARAReportWorkbookRenderer().render(
        view, ROOT / "references" / "HARA_Template_AI_20260327.xlsx", output, schema,
    )
    workbook = load_workbook(output)
    try:
        hara = workbook["04_HARA"]
        detail = workbook["04A_Scenario Detail"]
        assert hara["X6"].value == view.rows[0].remark
        assert hara["X7"].value == view.rows[1].remark
        headers = {detail.cell(4, column).value: column for column in range(1, detail.max_column + 1)}
        assert detail.cell(5, headers["Physical Inputs"]).value == view.scenario_details[0].physical_inputs
        assert detail.cell(6, headers["Driver Branch"]).value == view.scenario_details[1].driver_branch
        assert all(
            merged.max_row < 6 or merged.min_col > 24
            for merged in hara.merged_cells.ranges
        )
    finally:
        workbook.close()


@pytest.mark.parametrize(("field", "change"), (
    ("severity", {"result": "S3"}),
    ("exposure", {"status": "PENDING_INPUT"}),
    ("controllability", {"result": "C1"}),
    ("asil", {"status": "PENDING_UPSTREAM_RISK_VALUE"}),
))
def test_scored_trace_field_conflict_blocks_report_projection(field, change):
    state = _state(("车辆可能与行人碰撞",) * 3)
    state.scenarios = state.scenarios[:1]
    state.risk_results = state.risk_results[:1]
    risk = state.risk_results[0]
    values = {"severity": "S2", "exposure": "E3", "controllability": "C2", "asil": "B"}
    for name, value in values.items():
        setattr(risk, name, EvidenceValue(value, ReviewStatus.FINALIZED, rule_version="method-v1"))
    trace = {
        "malfunction_id": "MF-1", "scenario_id": "SCN-1",
        "risk_scoring_invoked": True,
        **{name: {"status": "FINALIZED", "result": value} for name, value in values.items()},
    }
    trace[field].update(change)

    with pytest.raises(ValueError, match=f"disagrees with {field} state"):
        HARAReportProjectionService(load_report_schema()).project(
            state, _method(), risk_trace={"assessments": [trace]},
        )


def test_supplied_trace_requires_every_risk_pair_and_rejects_duplicate_rows():
    state = _state(("车辆可能与行人碰撞",) * 3)
    service = HARAReportProjectionService(load_report_schema())
    one = {
        "malfunction_id": "MF-1", "scenario_id": "SCN-1",
        "risk_scoring_invoked": False,
    }

    with pytest.raises(ValueError, match="lacks committed risk pairs"):
        service.project(state, _method(), risk_trace={"assessments": [one]})
    with pytest.raises(ValueError, match="duplicate risk pair"):
        service.project(state, _method(), risk_trace={"assessments": [one, dict(one)]})


def test_scored_trace_requires_each_field_status_and_result():
    state = _state(("车辆可能与行人碰撞",) * 3)
    state.scenarios = state.scenarios[:1]
    state.risk_results = state.risk_results[:1]
    trace = {
        "malfunction_id": "MF-1", "scenario_id": "SCN-1",
        "risk_scoring_invoked": True,
        "severity": {"status": "PENDING_INPUT", "result": None},
        "exposure": {"status": "PENDING_INPUT", "result": None},
        "controllability": {"status": "PENDING_INPUT", "result": None},
    }

    with pytest.raises(ValueError, match="lacks asil status/result"):
        HARAReportProjectionService(load_report_schema()).project(
            state, _method(), risk_trace={"assessments": [trace]},
        )


def test_unscored_physical_values_are_labeled_as_scenario_candidates():
    scenario = _scenario(1)
    scenario.facts.update({"ego_speed_kph": 7, "relative_distance_m": 3, "ttc_s": 1.54})
    mapper = EngineeringReportTextMapper()

    candidate_text = mapper.physical_inputs(scenario, {"risk_scoring_invoked": False})
    assert candidate_text.startswith("场景候选（未核实用于评分）：")
    assert "自车速度：7 km/h" in candidate_text
    accepted_text = mapper.physical_inputs(scenario, {
        "risk_scoring_invoked": True,
        "hazardous_event_risk_context": {
            "ego_speed_kph": {"status": "AVAILABLE", "value": 7},
        },
    })
    assert "自车速度：7 km/h" in accepted_text
    assert "相对距离：3 m" not in accepted_text
    assert "TTC：1.54 s" not in accepted_text


def test_canonical_renderer_reads_current_run_trace_and_writes_real_reference(tmp_path, monkeypatch):
    from hara_agent.services.reporting import HARAExcelRenderer

    state = _state(("车辆可能与行人碰撞",) * 3)
    state.scenarios = state.scenarios[:1]
    state.risk_results = state.risk_results[:1]
    review_root = tmp_path / "review"
    trace_path = review_root / state.run_id / "risk_execution_trace.json"
    trace_path.parent.mkdir(parents=True)
    trace_path.write_text(json.dumps({
        "run_id": state.run_id,
        "method_contract_hash": "method-hash",
        "assessments": [{
            "malfunction_id": "MF-1", "scenario_id": "SCN-1",
            "risk_scoring_invoked": True,
            "severity": {"status": "PENDING_INPUT", "result": None},
            "exposure": {"status": "PENDING_INPUT", "result": None},
            "controllability": {
                "status": "PENDING_INPUT", "result": None,
                "decision_tree_stage": "TTC", "rule_ids": ["C-TTC-01"],
                "unknown_override_policy": "UNSPECIFIED",
            },
            "asil": {"status": "PENDING_UPSTREAM_RISK_VALUE", "result": None},
            "hazardous_event_risk_context": {
                "ego_speed_kph": {"status": "AVAILABLE", "value": 7},
            },
        }],
    }), encoding="utf-8")
    monkeypatch.setenv("HARA_REVIEW_ARTIFACT_DIR", str(review_root))
    output = tmp_path / "report.xlsx"
    HARAExcelRenderer(
        report_schema=load_report_schema(), method_contract=_method(),
    ).render(
        state, ROOT / "references" / "HARA_Template_AI_20260327.xlsx",
        output, draft=True,
    )

    workbook = load_workbook(output, read_only=True)
    try:
        detail = workbook["04A_Scenario Detail"]
        detail_headers = {
            cell.value: cell.column for cell in detail[4] if cell.value is not None
        }
        assert "C-TTC-01" in detail.cell(5, detail_headers["C Decision"]).value
        assert "自车速度：7 km/h" in detail.cell(5, detail_headers["Physical Inputs"]).value
        audit = workbook["99_Audit"]
        audit_headers = {cell.value: cell.column for cell in audit[4] if cell.value is not None}
        assert audit.cell(5, audit_headers["Risk execution trace"]).value == str(trace_path)
    finally:
        workbook.close()


def test_canonical_renderer_blocks_scored_state_without_trace(tmp_path, monkeypatch):
    from hara_agent.services.reporting import HARAExcelRenderer

    state = _state(("车辆可能与行人碰撞",) * 3)
    state.scenarios = state.scenarios[:1]
    state.risk_results = state.risk_results[:1]
    monkeypatch.setenv("HARA_REVIEW_ARTIFACT_DIR", str(tmp_path / "empty-review"))
    output = tmp_path / "report.xlsx"

    with pytest.raises(ValueError, match="Risk execution trace is required"):
        HARAExcelRenderer(
            report_schema=load_report_schema(), method_contract=_method(),
        ).render(
            state, ROOT / "references" / "HARA_Template_AI_20260327.xlsx",
            output, draft=True,
        )
    assert not output.exists()
