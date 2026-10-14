"""Fixed-population P7 replay using the existing production services (zero Provider).

The source checkpoint and its artifacts are read-only. This harness creates a
new run, traces input consumption, and prepares a bounded live-validation plan.
It contains no S/E/C or ASIL business rules.
"""
import argparse
from collections import Counter, defaultdict
from copy import deepcopy
from dataclasses import asdict
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from hara_agent.method_sources import MethodSourceResolver
from hara_agent.models import ItemDefinitionFacts, ScenarioFeasibilityAssessment, ReviewStatus
from hara_agent.services.analysis import MethodContractASILService, MethodRiskFactBindingService, StructuredRiskScoringService
from hara_agent.services.analysis.scenario_physics import causal_physics_conflicts, select_ego_speed_from_policy, restore_template_ego_direction
from hara_agent.services.analysis.driver_configuration_service import DriverConfigurationBrancher
from hara_agent.services.analysis.scenario_equivalence import signature, scenario_conditions
from hara_agent.services.semantic.scenario_agent import ScenarioFeasibilityAgent
from hara_agent.workflow import HARAState, ReviewArtifactWriter
from hara_agent.workflow.checkpoints import CheckpointRepository
from hara_agent.workflow.nodes.scoring import score_structured_scenarios
from hara_agent.workflow.risk_rescoring import clean_risk_stage
from hara_agent.services.reporting.offline_rebuild import OfflineReportRebuilder


def write(path, payload):
    path.write_text(json.dumps(HARAState._serialize(payload), ensure_ascii=False, indent=2, default=str) + "\n")


def run(args):
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    source_hash = hashlib.sha256(args.checkpoint.read_bytes()).hexdigest()
    raw = json.loads(args.checkpoint.read_text())
    state = HARAState.read_committed(deepcopy(raw))
    method = MethodSourceResolver().resolve(template_path=None, baseline_manifest_path=args.baseline,
        report_template_path=args.template).method
    state.run_id = args.run_id
    state.method_contract.update(template_hash=method.metadata['template_hash'],
        method_source_hash=method.metadata['method_source_hash'], compiler_version=method.compiler_version)
    old_risks = {(r.malfunction_id, r.scenario_id): r for r in state.risk_results}
    old_goals = deepcopy(state.safety_goals)
    clean_risk_stage(state)
    catalog = {a['atom_id']: a for a in method.metadata['scenario_atom_catalog']}
    policy = method.metadata['project_analysis_policy']
    brancher = DriverConfigurationBrancher(policy_id=policy['policy_id'], policy=policy)
    changed = []
    for index, candidate in enumerate(state.scenarios):
        original = deepcopy(candidate.facts)
        recovered = restore_template_ego_direction(candidate, method.scenario_model.scenario_method.fm_template_catalog.templates)
        if recovered:
            candidate.facts['ego_longitudinal_direction'], candidate.fact_provenance['ego_longitudinal_direction'] = recovered
        for binding in candidate.facts.get('method_scenario_dimensions', {}).values():
            atom = catalog.get(binding.get('atom_id'))
            if atom:
                binding['method_semantics'] = deepcopy(atom.get('physical_semantics', atom.get('method_semantics', {})))
        # Only retire the historical policy-selected normal speed, never a
        # measured/project fact. Recompute it with current time semantics.
        meta = candidate.fact_provenance.get('ego_speed_kph', {})
        if meta.get('selection_basis') == 'GOVERNED_CLOSED_UPPER_BOUND':
            candidate.facts.pop('ego_speed_kph', None)
            candidate.fact_provenance.pop('ego_speed_kph', None)
        mf_id = candidate.analysis_instance.get('malfunction_id', '')
        speed = select_ego_speed_from_policy(candidate, policy=policy, malfunction_id=mf_id)
        if speed:
            candidate.facts['ego_speed_kph'], candidate.fact_provenance['ego_speed_kph'] = speed
        candidate = brancher.apply_location_convention(candidate)
        state.scenarios[index] = candidate
        differences = [key for key in set(original) | set(candidate.facts) if original.get(key) != candidate.facts.get(key)]
        if differences:
            changed.append({'scenario_id': candidate.scenario_id, 'fields': sorted(differences)})
    by_id = {c.scenario_id: c for c in state.scenarios}
    rejected = []
    for assessment in state.item_definition['scenario_assessments']:
        c = by_id.get(assessment['scenario_id'])
        if c is None or not assessment.get('causally_relevant'):
            continue
        conflicts = causal_physics_conflicts(c, assessment)
        if conflicts:
            rejected.append({'malfunction_id': assessment['malfunction_id'], 'scenario_id': c.scenario_id,
                'reason_codes': conflicts, 'original_hazardous_event': assessment.get('hazardous_event')})
            assessment['status'] = 'PENDING'
            if assessment.get('causal_assessment'):
                assessment['causal_assessment']['review_status'] = 'PENDING'
            assessment['rationale'] = 'P7 physical consistency requires causal revalidation: ' + ','.join(conflicts)
    state.record('p7_fixed_population_replay', source_run_id=raw['run_id'], source_checkpoint_sha256=source_hash,
        policy_id=policy['policy_id'], semantic_revalidation_required=rejected)
    review_root = out / 'review'
    writer = ReviewArtifactWriter(state.run_id, review_root)
    for function in state.functions:
        writer.record_function(function)
    for mf in state.malfunctions:
        writer.record_malfunction(mf)
    mf_by_scenario = {a['scenario_id']: a['malfunction_id'] for a in state.item_definition['scenario_assessments']}
    for c in state.scenarios:
        writer.record_scenario_candidate(c, generated_for_malfunction_ids=[mf_by_scenario.get(c.scenario_id, '')])
    for a in state.item_definition['scenario_assessments']:
        writer.record_scenario_feasibility(a)
    state = score_structured_scenarios(state, StructuredRiskScoringService(method), MethodContractASILService(method),
        risk_fact_binding=MethodRiskFactBindingService(method), review_artifact_writer=writer)
    # Retain existing SG associations as historical drafts; no new SG business
    # generation is performed in P7 (G-1).
    state.safety_goals = old_goals
    for risk in state.risk_results:
        previous = old_risks.get((risk.malfunction_id, risk.scenario_id))
        if previous:
            risk.safety_goal_id = previous.safety_goal_id
    repository = CheckpointRepository(out / 'checkpoints')
    checkpoint = repository.save(state)
    event = next(e for e in reversed(state.audit_trail) if e['event'] == 'structured_risk_scoring_completed')
    write(out / 'input_consumption.json', {'bindings': event['risk_fact_binding_audits'],
        'calculation_inputs': event['risk_calculation_inputs']})
    def counts(risks):
        return {field: dict(Counter((getattr(risk, field).value or 'Pending') for risk in risks))
                for field in ('severity', 'exposure', 'controllability', 'asil')}
    summary = {'source_run_id': raw['run_id'], 'run_id': state.run_id,
        'source_checkpoint_sha256': source_hash, 'method_hash': method.metadata['method_source_hash'],
        'compiler_version': method.compiler_version, 'provider_calls': 0,
        'fixed_candidate_count': len(state.item_definition['scenario_assessments']),
        'before': counts(old_risks.values()), 'after': counts(state.risk_results),
        'new_causal_pending': len(rejected), 'fact_changes': len(changed),
        'input_binding_counts': dict(Counter(t for a in event['risk_fact_binding_audits'] for t in a['bound_fact_types'])),
        'binding_conflicts': sum(bool(a['source_conflicts']) for a in event['risk_fact_binding_audits']),
        'all_dimensions_finalized': sum(all(getattr(r,f).status is ReviewStatus.FINALIZED for f in ('severity','exposure','controllability','asil')) for r in state.risk_results),
        'asil_decision_finalized': sum(r.asil.status is ReviewStatus.FINALIZED for r in state.risk_results)}
    write(out / 'fixed_population_summary.json', summary)
    write(out / 'causal_revalidation_required.json', rejected)
    write(out / 'fact_changes.json', changed)
    workbook = load_workbook(args.source_report, read_only=True, data_only=True)
    mapping = {r[1]: (r[2], r[3]) for r in workbook['99_Audit'].iter_rows(min_row=5, values_only=True) if r[1]}
    workbook.close()
    write(out / 'report_identity_map.json', mapping)
    # Representative historical cases, one companion driver branch each where
    # available. Preserve individual identities and independent judgments.
    targets = ['HARA_250','HARA_274','HARA_020','HARA_457','HARA_064','HARA_621','HARA_039','HARA_002']
    candidates_by_conditions = defaultdict(list)
    for c in state.scenarios:
        candidates_by_conditions[(mf_by_scenario.get(c.scenario_id), signature(scenario_conditions(c, without_driver=True)))].append(c)
    selected = defaultdict(dict)
    for report_id in targets:
        mf, sid = mapping[report_id]
        sc = by_id[sid]
        key = (mf, signature(scenario_conditions(sc, without_driver=True)))
        for c in candidates_by_conditions[key][:2]:
            selected[mf][c.scenario_id] = c
    functions = {f['function_id']: f for f in state.functions}
    malfunctions = {m['malfunction_id']: m for m in state.malfunctions}
    plan = {'project': raw['item_definition']['typed'], 'source_checkpoint_sha256': source_hash,
            'method_hash': method.metadata['method_source_hash'], 'cases': []}
    for mf, cases in selected.items():
        plan['cases'].append({'malfunction': malfunctions[mf], 'function': functions[malfunctions[mf]['function_id']],
            'candidates': [asdict(c) for c in cases.values()], 'generation': {'source': 'P7_FIXED_HISTORICAL_SAMPLE'}})
    write(out / 'sample_plan.json', plan)
    result = OfflineReportRebuilder().rebuild(checkpoint_path=checkpoint, method_baseline_path=args.baseline,
        report_style_template_path=args.template, output_path=out / 'HARA_P7_Revised_Draft_20261014.xlsx',
        review_root=review_root, audit_output_dir=out / 'report_audit')
    print(json.dumps({'summary': summary, 'report': str(result),
                      'sample_instances': sum(len(c['candidates']) for c in plan['cases'])}, ensure_ascii=False, default=str))
    assert hashlib.sha256(args.checkpoint.read_bytes()).hexdigest() == source_hash


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--checkpoint', type=Path, default=Path('runtime/agent/full-p6-20261012-002.checkpoint.json'))
    parser.add_argument('--source-report', type=Path, default=Path('output/HARA_full_p6_20261012_002.xlsx'))
    parser.add_argument('--output', type=Path, default=Path('output/p7_closure'))
    parser.add_argument('--run-id', default='p7-offline-20261014-001')
    parser.add_argument('--baseline', type=Path, default=Path('method_assets/fusa_baseline_v1/manifest.yaml'))
    parser.add_argument('--template', type=Path, default=Path('references/HARA_Template_AI_20260327.xlsx'))
    run(parser.parse_args())
