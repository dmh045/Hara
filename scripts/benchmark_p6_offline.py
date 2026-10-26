"""Replay P6 candidate generation and E previews without creating a Provider.

This is a benchmark of the production services, not a new analysis runtime.
New candidates have NOT passed fresh causal validation; previews are reported
separately from committed historical risk ratings.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import time
from types import SimpleNamespace

from openpyxl import load_workbook

from hara_agent.method_sources import MethodSourceResolver
from hara_agent.application import HARAApplication
from hara_agent.config import RunConfig
from hara_agent.workflow import HARAState
from hara_agent.services.extraction import DocumentReader
from hara_agent.services.semantic.item_definition_agent import ItemDefinitionNormalizer
from hara_agent.models import ItemDefinitionFacts, MalfunctionCandidate, ReviewStatus, SourceRef
from hara_agent.services.analysis import MethodScenarioCandidateService, ProjectFactResolver
from hara_agent.services.analysis.exposure_input_readiness_service import ExposureInputReadinessService, READY_STATUSES
from hara_agent.services.analysis.malfunction_situation_selection import selection_function
from hara_agent.services.analysis.scenario_equivalence import driver_families, scenario_conditions, signature


def historical_audit(saved: dict, review: Path, corrected_report: Path | None) -> dict:
    """Measure identities and structured similarities without declaring new causality."""
    files = [review / name for name in ("scenario_candidates.jsonl", "scenario_binding_gaps.json",
                                        "malfunction_situation_selection.json", "risk_execution_trace.json")]
    manifest = []
    for path in files + ([corrected_report] if corrected_report else []):
        with path.open('rb') as stream:
            manifest.append({'path': str(path), 'sha256': hashlib.file_digest(stream, 'sha256').hexdigest()})
    candidates = {}
    with files[0].open() as stream:
        for line in stream:
            row = json.loads(line)
            candidates[row['scenario_id']] = row
    parents = [row for row in candidates.values() if not row.get('analysis_instance')]
    bindings = Counter(binding.get('resolution_status', 'MISSING')
                       for row in parents for binding in row['facts'].get('method_scenario_dimensions', {}).values())
    gaps = json.loads(files[1].read_text())
    selection = json.loads(files[2].read_text())
    trace = json.loads(files[3].read_text())
    pairs, exact, driver, physical, shared_worlds, mechanism_variants = [], Counter(), Counter(), Counter(), defaultdict(set), defaultdict(set)
    motion = Counter()
    physical_keys = ('object_type', 'object_position', 'road_user_type', 'collision_type', 'relative_distance',
                     'relative_distance_m', 'object_speed_kph', 'ego_speed_constraint', 'ego_longitudinal_direction',
                     'object_longitudinal_direction', 'harm_mechanism')
    for assessment in saved['item_definition']['scenario_assessments']:
        mf, sid = assessment['malfunction_id'], assessment['scenario_id']
        candidate = candidates[sid]
        facts = dict(candidate['facts'])
        # Historical v3 compact review omitted these explicit branch fields.
        # Recover the saved position from branch metadata, never from prose.
        branch = candidate.get('analysis_instance', {}).get('driver_configuration_branch', {})
        if branch.get('driver_position'):
            facts['driver_position'] = branch['driver_position']
        obj = SimpleNamespace(scenario_id=sid, facts=facts, operating_mode=candidate['operating_mode'],
                              fact_provenance=candidate.get('fact_provenance', {}))
        pairs.append((mf, sid))
        exact[(mf, signature(scenario_conditions(obj)))] += 1
        driver[(mf, signature(scenario_conditions(obj, without_driver=True)))] += 1
        physical[(mf, signature({k: obj.facts.get(k) for k in physical_keys}))] += 1
        shared_worlds[signature(scenario_conditions(obj))].add(mf)
        mechanism_variants[mf].add(signature({k: obj.facts.get(k) for k in physical_keys}))
        motion[str(obj.facts.get('ego_longitudinal_direction'))] += 1
    workbook = None
    if corrected_report:
        book = load_workbook(corrected_report, read_only=True, data_only=True)
        try:
            workbook = {}
            for name in ('04_HARA', '04A_Scenario Detail', '99_Audit'):
                sheet = book[name]
                headers = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
                identity = headers.index('HARA-ID' if name == '04_HARA' else 'Scenario ID')
                workbook[name] = {'data_rows': sum(bool(row[identity]) for row in sheet.iter_rows(min_row=5, values_only=True))}
        finally:
            book.close()
    return {
        'source_manifest': manifest,
        'parents': {'count': len(parents), 'empty_atom_sets': sum(not p['facts'].get('scenario_atom_ids') for p in parents),
                    'binding_statuses': dict(bindings)},
        'binding_gap_reasons': dict(Counter(row['reason'] for row in gaps['gaps'])),
        'dimension_coverage': gaps['scenario_binding_coverage']['dimensions'],
        'selection_counts': {key: sum(row.get(key, 0) for row in selection['per_malfunction'])
                             for key in ('base_candidate_count', 'include_count', 'exclude_count', 'unresolved_passthrough_count')},
        'risk_eligibility': dict(Counter(row['risk_eligibility']['status'] for row in trace['assessments'])),
        'redundancy': {
            'identity_duplicate_pairs': len(pairs) - len(set(pairs)),
            'same_structured_inputs_groups_within_malfunction': len(exact),
            'same_structured_inputs_redundant_pairs': sum(n - 1 for n in exact.values()),
            'same_non_driver_inputs_groups': len(driver),
            'potential_driver_reuse_pairs': sum(n - 1 for n in driver.values()),
            'same_physical_inputs_ignoring_environment_groups': len(physical),
            'potential_common_behavior_pairs_across_environment': sum(n - 1 for n in physical.values()),
            'physical_variants_per_malfunction': {mf: len(values) for mf, values in mechanism_variants.items()},
            'identical_worlds_shared_by_multiple_malfunctions': sum(len(values) > 1 for values in shared_worlds.values()),
            'ego_direction_distribution': dict(motion),
            'proven_causal_reuse': 0, 'proven_complete_risk_equivalence': 0,
            'note': 'Groups use persisted compact structured inputs plus explicit driver branch identity; they are audit leads, not accepted equivalence. Historical atoms are empty and dependency metadata is absent. No cross-malfunction causal reuse.',
        },
        'corrected_report': workbook,
    }


def run(checkpoint: Path, method_baseline: Path, report_template: Path, output: Path,
        corrected_report: Path | None = None, item: Path | None = None, *, generation_only: bool = False,
        operating_mode: str | None = None) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    with checkpoint.open('rb') as stream:
        before_hash = hashlib.file_digest(stream, 'sha256').hexdigest()
    saved = json.loads(checkpoint.read_text())
    method = MethodSourceResolver().resolve(template_path=None, baseline_manifest_path=method_baseline,
                                          report_template_path=report_template).method
    # Replay deliberately compares historical inputs with the active method.
    # Both identities are reported; no historical risk grade is overwritten.
    typed = dict(saved['item_definition']['typed'])
    # Derived facts from old scenario identities do not become new project input.
    typed['risk_facts'] = [fact for fact in typed['risk_facts'] if 'scenario_id' not in fact.get('context', {})]
    fact_ids = {fact['fact_id'] for fact in typed['risk_facts']}
    typed['method_risk_fact_bindings'] = [b for b in typed['method_risk_fact_bindings'] if b['source_fact_id'] in fact_ids]
    project = ItemDefinitionFacts.from_dict(typed)
    if item is not None:
        document = DocumentReader().read(item)
        project = ItemDefinitionNormalizer.normalize_source_scope(project, document.blocks, document.source_id)
    service = MethodScenarioCandidateService(method)
    mode = operating_mode or next((s['operating_mode'] for s in saved['scenarios'] if s.get('operating_mode')), None)
    if not mode:
        raise ValueError('A checkpoint without generated scenarios requires --operating-mode')
    speed = ProjectFactResolver().resolve_speed_context(project, mode)
    functions = {f['function_id']: selection_function(f) for f in saved['functions']}
    parents, parent_audit = service.generate(project_facts=project, operating_mode=mode,
                                           speed_resolution=speed, functions=list(functions.values()))
    if item is not None:
        # Exercise the ordinary application boundary with the same saved
        # extraction and explicit source normalization; no client is created.
        state = HARAState(run_id="offline-coverage", item_definition={"typed": asdict(project)},
                          functions=saved['functions'])
        application = HARAApplication(RunConfig(item_path=item, template_path=None, run_dir=output,
            output_path=output / 'unused.xlsx',
            method_baseline_path=method_baseline, report_template_path=report_template,
            operating_mode=mode, allow_draft=True), None)
        parents, parent_audit = application.prepare_scenario_candidates(state, service)
    readiness = ExposureInputReadinessService(method)
    counts, e_reasons, e_values, ready_e_values, dimensional = Counter(), Counter(), Counter(), Counter(), Counter()
    source_locations, source_roads, where_atoms, road_atoms = Counter(), Counter(), Counter(), Counter()
    per_malfunction, candidate_records = [], []
    start = time.monotonic()
    for index, raw in enumerate(saved['malfunctions'], 1):
        malfunction = MalfunctionCandidate(**{**raw, 'status': ReviewStatus(raw['status']),
                                             'sources': [SourceRef(**source) for source in raw['sources']]})
        children, audit = service.generate_for_malfunction(
            malfunction, parents, project_facts=project, function=functions.get(malfunction.function_id),
        )
        generation = audit['atom_generation']
        counts['raw_generated_before_driver'] += generation['raw_candidate_count']
        counts['unique_conditions_before_driver'] += generation['unique_condition_count']
        counts['coalesced_candidates'] += generation['coalesced_candidate_count']
        context_audits = [context for entry in generation['parents']
                          for context in entry.get('source_contexts', [entry])]
        counts['planned_pool_before_driver'] += sum(a.get('planned_pool_candidate_count', 0) for a in context_audits)
        counts['deferred_context_variant_combinations'] += sum(len(a.get('deferred_context_variants', [])) for a in context_audits)
        counts['candidate_scenarios_after_driver'] += len(children)
        families = driver_families(children)
        counts['driver_families'] += len(families)
        counts['potential_driver_reuse_requires_dependency_proof'] += len(children) - len(families)
        for child in children:
            counts['collision_branches'] += len(child.facts.get('collision_branches', []))
            counts['motion_models'] += bool(child.facts.get('fault_motion_model'))
            counts['impact_speeds'] += child.facts.get('ego_impact_speed_kph') is not None
            atoms = child.facts.get('scenario_atom_ids', [])
            source_conditions = child.facts.get('source_operating_conditions', {})
            source_locations[str(source_conditions.get('operating_scenario', ''))] += 1
            source_roads[str(source_conditions.get('road_surface_conditions', ''))] += 1
            bindings = child.facts.get('method_scenario_dimensions', {})
            where_atoms[str(bindings.get('WHERE', {}).get('atom_id', ''))] += 1
            road_atoms[str(bindings.get('ROAD', {}).get('atom_id', ''))] += 1
            counts['nonempty_atom_sets'] += bool(atoms)
            result = readiness.assess({**child.facts, 'component_category': malfunction.component_category})
            counts['e_preview_ready'] += result['status'] in READY_STATUSES
            if result['status'] in READY_STATUSES:
                ready_e_values[result['baseline_exposure']['value']] += 1
            counts['method_constraint_validated'] += child.analysis_instance.get('validation_status') == 'VALIDATED'
            e_reasons[result['reason_code']] += 1
            e_values[result['baseline_exposure']['value']] += 1
            for dimension, binding in child.facts.get('method_scenario_dimensions', {}).items():
                dimensional[dimension + ':' + binding.get('resolution_status', 'MISSING')] += 1
            candidate_records.append({
                'malfunction_id': malfunction.malfunction_id, 'scenario_id': child.scenario_id,
                'scenario_atom_ids': atoms, 'condition_signature': signature(scenario_conditions(child)),
                'ego_direction': child.facts.get('ego_longitudinal_direction'),
                'ego_speed_constraint': child.facts.get('ego_speed_constraint'),
                'driver_position': child.facts.get('driver_position'),
                'e_preview': result, 'causal_status': 'NOT_REVALIDATED',
                'source_operating_conditions': source_conditions,
                'physical_input_production': child.analysis_instance.get('physical_input_production', {}),
                'collision_branches': child.facts.get('collision_branches', []),
                'source_candidate': asdict(child),
                'method_conditions': {dim: {key: value.get(key) for key in ('atom_id', 'atom_label', 'resolution_status')}
                                      for dim, value in bindings.items()},
                'validation_status': child.analysis_instance.get('validation_status'),
            })
        per_malfunction.append({'malfunction_id': malfunction.malfunction_id, 'candidate_count': len(children),
                               'generation': generation, 'driver_configuration': audit.get('driver_configuration', {})})
        print(f'[P6 offline] {index}/{len(saved["malfunctions"])} {malfunction.malfunction_id}: {len(children)} candidates', flush=True)
    old_audit = ({'status': 'NOT_REQUESTED_GENERATION_ONLY'} if generation_only else
        historical_audit(saved, checkpoint.parent.parent / 'review' / saved['run_id'], corrected_report))
    old_assessments = saved['item_definition'].get('scenario_assessments', [])
    latest = {a['malfunction_id']: a for a in saved['audit_trail'] if a.get('event') == 'scenario_feasibility_assessed'}
    historical_ratings = {field: dict(Counter(r[field]['status'] for r in saved['risk_results']))
                          for field in ('severity', 'exposure', 'controllability', 'asil')}
    result = {
        'source_run_id': saved['run_id'], 'provider_calls': 0, 'retries': 0,
        'input_tokens': 0, 'output_tokens': 0, 'elapsed_seconds': round(time.monotonic() - start, 3),
        'method_hash': method.metadata['method_source_hash'],
        'source_method_hash': saved['method_contract']['method_source_hash'],
        'historical': {'scenario_malfunction_pairs': len(old_assessments),
                       'risk_records': len(saved['risk_results']), 'ratings': historical_ratings,
                       'causal_provider_calls': sum(a.get('actual_llm_calls', 0) for a in latest.values()),
                       'causal_retries': sum(a.get('retry_calls', 0) for a in latest.values()),
                       'causal_input_tokens': sum(a.get('prompt_tokens_total', 0) for a in latest.values()),
                       'causal_output_tokens': sum(a.get('completion_tokens_total', 0) for a in latest.values())},
        'candidate_preview': {**counts, 'e_preview_reasons': dict(e_reasons), 'all_trial_e_values_including_blocked': dict(e_values),
                              'ready_e_preview_values': dict(ready_e_values),
                              'dimension_statuses': dict(dimensional), 'generated_parent_count': len(parents)},
        'source_coverage': {'locations': dict(source_locations), 'road_contexts': dict(source_roads),
                            'scope_evidence': project.odd_scope_evidence,
                            'permitted_road_types': list(project.odd_road_types),
                            'road_type_coverage_status': 'SOURCE_SCOPE_PRESERVED; RISK_EQUIVALENCE_NOT_PROVEN',
                            'WHERE': dict(where_atoms), 'ROAD': dict(road_atoms),
                            'policy': 'Joint coverage of primary variants, source contexts and WHERE/ROAD alternatives; unselected combinations remain unevaluated. Not exhaustive joint ODD coverage.'},
        'new_committed_risk_records': 0, 'new_causal_provider_validation': 'NOT_RUN',
        'new_s_c_asil': 'NOT_REVALIDATED; historical ratings unchanged',
        'candidate_generation_provider_calls': 0,
        'source_checkpoint_sha256': before_hash,
    }
    with checkpoint.open('rb') as stream:
        assert hashlib.file_digest(stream, 'sha256').hexdigest() == before_hash
    for filename, payload in [('summary.json', result), ('candidate_previews.json', candidate_records),
                              ('generation_audit.json', per_malfunction), ('parent_binding_audit.json', parent_audit),
                              ('historical_audit.json', old_audit)]:
        (output / filename).write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--checkpoint', type=Path, required=True)
    parser.add_argument('--method-baseline', type=Path, default=Path('method_assets/fusa_baseline_v1/manifest.yaml'))
    parser.add_argument('--report-template', type=Path, default=Path('references/HARA_Template_AI_20260327.xlsx'))
    parser.add_argument('--output', type=Path, default=Path('output/p6'))
    parser.add_argument('--corrected-report', type=Path)
    parser.add_argument('--item', type=Path, help='Reapply production source-scope normalization to the original Item')
    parser.add_argument('--generation-only', action='store_true', help='Compare production candidate generation without requiring historical review sidecars')
    parser.add_argument('--operating-mode', help='Explicit scope for a pre-generation checkpoint')
    args = parser.parse_args()
    run(args.checkpoint, args.method_baseline, args.report_template, args.output, args.corrected_report, args.item,
        generation_only=args.generation_only, operating_mode=args.operating_mode)
